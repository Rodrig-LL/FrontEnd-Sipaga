"""Workbook laporan validasi impor DPP (terpisah dari file pembukuan)."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from dpp_parser import KOLOM_ANGKA

F = "Arial"
JUDUL = Font(name=F, size=13, bold=True)
TEBAL = Font(name=F, size=10, bold=True)
BIASA = Font(name=F, size=10)
PUTIH = Font(name=F, size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
TIPIS = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
RP = '#,##0;(#,##0);"-"'
WARNA_TINGKAT = {"TINGGI": "F8D7DA", "SEDANG": "FFF3CD", "INFO": "E2EFDA"}

LABEL_ANGKA = {
    "g_pokok": "G.POKOK", "t_istri": "T.ISTRI/SUAMI", "t_anak": "T.ANAK", "g_bruto": "G.BRUTO",
    "t_struktural": "T.STRUKTURAL", "t_fungsional": "T.FUNGSIONAL", "t_umum": "T.UMUM",
    "t_beras": "T.PGN/BERAS", "t_kowan": "T.KOWAN", "t_sandi": "T.SANDI/KOMPEN",
    "t_babinsa": "T.BABINSA", "t_papua": "T.PAPUA", "t_pencil": "T.PENCIL",
    "t_terluar": "T.P.TERLUAR", "t_terampil": "T.TERAMPIL", "t_lainnya": "T.LAINNYA", "tpp": "TPP",
    "t_pph": "T.PPh", "pembulatan": "PEMBULATAN", "jml_kotor": "JML KOTOR", "p_beras": "P.BERAS",
    "p_pensiun": "P.PENSIUN", "p_bpjs": "P.BPJS", "p_tht": "P.THT", "p_bpjs_lain": "P.BPJS LAIN",
    "p_sewa_rumah": "SEWA RUMAH", "p_utang": "UTANG", "p_pph21": "PPh PS.21",
    "p_pengembalian": "PENGEMBALIAN", "p_tgr": "TGR", "p_lainnya": "POT LAINNYA",
    "jml_potongan": "JML POTONGAN", "bersih": "BERSIH GAJI", "lauk_pauk": "LAUK PAUK",
    "dibayarkan": "JML DIBAYARKAN",
}


def _header(ws, row, judul, lebar):
    for i, (j, w) in enumerate(zip(judul, lebar), 1):
        c = ws.cell(row, i, j)
        c.font, c.fill, c.border = PUTIH, HEAD_FILL, TIPIS
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def _isi(ws, row, nilai, fmt_kolom=()):
    for i, v in enumerate(nilai, 1):
        c = ws.cell(row, i, v)
        c.font, c.border = BIASA, TIPIS
        if i in fmt_kolom:
            c.number_format = RP


def buat_laporan(path, ctx):
    wb = Workbook()

    # ---------------------------------------------------------------- RINGKASAN
    ws = wb.active
    ws.title = "RINGKASAN"
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"LAPORAN VALIDASI IMPOR DPP GAJI - BULAN {ctx['bulan_baru']}"
    ws["A1"].font = JUDUL
    ws["A2"] = (f"Dibuat {datetime.now():%d-%m-%Y %H:%M} | Buku sumber: {ctx['buku_sumber']} "
                f"| Hasil: {ctx['buku_hasil']}")
    ws["A2"].font = Font(name=F, size=9, italic=True, color="666666")

    n_galat = sum(len(h["galat"]) for h in ctx["dpp_files"])

    r = 4
    ws.cell(r, 1, "A. PEMBACAAN PDF DPP").font = TEBAL
    r += 1
    _header(ws, r, ["File PDF", "Bulan", "Anggota", "Hitungan konsisten", "Lembar cocok subtotal",
                    "Gagal dibaca", "Jumlah dibayarkan (Rp)"], [34, 16, 11, 16, 16, 13, 22])
    r0 = r + 1
    for h in ctx["dpp_files"]:
        r += 1
        ok = sum(1 for a in h["anggota"] if not a["masalah"])
        _isi(ws, r, [h["file"], h["bulan"], len(h["anggota"]), f"{ok}/{len(h['anggota'])}",
                     f"{sum(l['cocok'] for l in h['lembar'])}/{len(h['lembar'])}",
                     len(h["galat"]),
                     sum(a["dibayarkan"] for a in h["anggota"])], fmt_kolom=(7,))
        if h["galat"]:
            ws.cell(r, 6).fill = PatternFill("solid", fgColor=WARNA_TINGKAT["TINGGI"])
            ws.cell(r, 6).font = TEBAL
    r += 1
    _isi(ws, r, ["TOTAL", "", f"=SUM(C{r0}:C{r-1})", "", "", f"=SUM(F{r0}:F{r-1})",
                 f"=SUM(G{r0}:G{r-1})"], fmt_kolom=(7,))
    for c in ws[r]:
        c.font = TEBAL

    # Kegagalan pembacaan tidak boleh tersembunyi: RINGKASAN harus ikut merah.
    if n_galat:
        r += 1
        c = ws.cell(r, 1, f"PERHATIAN: {n_galat} bagian PDF gagal dibaca. Data DPP di bagian itu "
                          "TIDAK ikut diproses dan tidak ada di sheet DATA DPP, sehingga angka di "
                          'atas belum lengkap. Rincian di sheet TEMUAN, kategori "Gagal membaca '
                          'data dari PDF".')
        c.font = Font(name=F, size=10, bold=True, color="9C0006")
        c.fill = PatternFill("solid", fgColor=WARNA_TINGKAT["TINGGI"])
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = TIPIS
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 32

    r += 2
    ws.cell(r, 1, "B. PENCOCOKAN DENGAN SHEET BUKU TNI").font = TEBAL
    s = ctx["statistik"]
    baris_stat = [
        ("Anggota di BUKU TNI", s["buku_total"], None),
        ("Berhasil dicocokkan & GAJI diperbarui", s["cocok"], None),
        ("  di antaranya gaji berubah dari bulan lalu", s["gaji_berubah"], None),
        ("Di buku tapi TIDAK ada di DPP (gaji tidak diperbarui)", s["tidak_di_dpp"], s["nilai_tidak_di_dpp"]),
        ("Ada di DPP tapi BELUM ada di buku", s["dpp_tanpa_buku"], s["nilai_dpp_tanpa_buku"]),
        ("Total GAJI hasil pencocokan (harus = jumlah DPP yang cocok)", None, s["total_update"]),
        ("Jumlah DPP yang cocok dengan buku", None, s["total_dpp_cocok"]),
        ("Label bulan yang diganti di seluruh workbook", s["label_bulan"], None),
    ]
    r += 1
    _header(ws, r, ["Uraian", "", "Jumlah", "", "", "", "Nilai (Rp)"],
            [34, 16, 11, 16, 16, 13, 22])
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    for u, n, v in baris_stat:
        r += 1
        _isi(ws, r, [u, "", n, "", "", "", v], fmt_kolom=(7,))
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    r += 1
    ws.cell(r, 1, "Rekonsiliasi").font = TEBAL
    ws.cell(r, 7, f'=IF(G{r-3}=G{r-2},"COCOK","SELISIH "&TEXT(G{r-3}-G{r-2},"#,##0"))').font = TEBAL

    r += 2
    ws.cell(r, 1, "C. TEMUAN PER TINGKAT").font = TEBAL
    r += 1
    _header(ws, r, ["Tingkat", "Arti", "Jumlah"], [34, 16, 11])
    arti = {"TINGGI": "Wajib dicek sebelum transfer", "SEDANG": "Perlu dikoreksi di buku",
            "INFO": "Untuk diketahui"}
    for tk in ("TINGGI", "SEDANG", "INFO"):
        r += 1
        _isi(ws, r, [tk, arti[tk], f'=COUNTIF(TEMUAN!A:A,"{tk}")'])
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=WARNA_TINGKAT[tk])
    ws.column_dimensions["B"].width = 28

    # ---------------------------------------------------------------- TEMUAN
    wt = wb.create_sheet("TEMUAN")
    _header(wt, 1, ["Tingkat", "Kategori", "Baris BUKU TNI", "NRP", "Nama", "Keterangan"],
            [10, 30, 10, 20, 34, 95])
    urut = {"TINGGI": 0, "SEDANG": 1, "INFO": 2}
    for i, t in enumerate(sorted(ctx["temuan"], key=lambda x: (urut[x["tingkat"]], x["kategori"],
                                                                 x["baris"] or 9999)), 2):
        _isi(wt, i, [t["tingkat"], t["kategori"], t["baris"], t["nrp"], t["nama"], t["detail"]])
        wt.cell(i, 1).fill = PatternFill("solid", fgColor=WARNA_TINGKAT[t["tingkat"]])
        wt.cell(i, 4).number_format = "@"
    wt.freeze_panes = "A2"
    wt.auto_filter.ref = f"A1:F{max(2, len(ctx['temuan']) + 1)}"

    # ---------------------------------------------------------------- PERUBAHAN GAJI
    wg = wb.create_sheet("PERUBAHAN GAJI")
    _header(wg, 1, ["Baris BUKU", "No", "Nama", "Pangkat", "NRP", "Gaji lama (Rp)", "Gaji baru (Rp)",
                    "Selisih (Rp)"], [10, 6, 34, 18, 20, 16, 16, 16])
    i = 1
    for b, d, _ in sorted(ctx["pasangan"], key=lambda x: x[0]["baris"]):
        if round(b["gaji_lama"]) != d["dibayarkan"]:
            i += 1
            _isi(wg, i, [b["baris"], b["no"], b["nama"], b["pangkat"], b["nrp_asli"],
                         b["gaji_lama"], d["dibayarkan"], f"=G{i}-F{i}"], fmt_kolom=(6, 7, 8))
    if i == 1:
        wg.cell(2, 1, "Tidak ada perubahan gaji dibanding isi buku sebelumnya.").font = BIASA
    else:
        _isi(wg, i + 1, ["", "", "TOTAL", "", "", f"=SUM(F2:F{i})", f"=SUM(G2:G{i})", f"=SUM(H2:H{i})"],
             fmt_kolom=(6, 7, 8))
        for c in wg[i + 1]:
            c.font = TEBAL
    wg.freeze_panes = "A2"

    # ---------------------------------------------------------------- DATA DPP
    wd = wb.create_sheet("DATA DPP")
    identitas = [("file", "File", 22), ("halaman", "Hal", 5), ("no_urut", "No", 5),
                 ("nama", "Nama", 32), ("nrp", "NRP", 18), ("korps", "Korps", 7),
                 ("pangkat", "Pangkat", 13), ("jabatan", "Jabatan", 28), ("tgl_lahir", "Tgl lahir", 11),
                 ("tmt_jab", "TMT jab", 11), ("jab_esln", "Jab/Esln", 13), ("sts_kawin", "Sts kawin", 9),
                 ("tmt_kgb", "TMT KGB", 11), ("mkg", "MKG", 7), ("npwp", "NPWP", 18),
                 ("rekening", "Rekening", 20)]
    judul = [j for _, j, _ in identitas] + [LABEL_ANGKA[k] for k in KOLOM_ANGKA] + \
            ["Baris BUKU TNI", "Status validasi"]
    lebar = [w for *_, w in identitas] + [13] * len(KOLOM_ANGKA) + [11, 40]
    _header(wd, 1, judul, lebar)
    n_id = len(identitas)
    peta_baris = {id(d): b["baris"] for b, d, _ in ctx["pasangan"]}
    semua = [a for h in ctx["dpp_files"] for a in h["anggota"]]
    for i, a in enumerate(semua, 2):
        _isi(wd, i, [a[k] for k, *_ in identitas] + [a[k] for k in KOLOM_ANGKA] +
             [peta_baris.get(id(a), "TIDAK ADA"), "; ".join(a["masalah"]) or "OK"],
             fmt_kolom=range(n_id + 1, n_id + len(KOLOM_ANGKA) + 1))
        for k in (5, 15, 16):
            wd.cell(i, k).number_format = "@"
    tot = len(semua) + 2
    wd.cell(tot, n_id, "TOTAL").font = TEBAL
    for j in range(n_id + 1, n_id + len(KOLOM_ANGKA) + 1):
        col = get_column_letter(j)
        c = wd.cell(tot, j, f"=SUM({col}2:{col}{tot-1})")
        c.font, c.number_format, c.border = TEBAL, RP, TIPIS
    wd.freeze_panes = wd.cell(2, 5)
    wd.auto_filter.ref = f"A1:{get_column_letter(len(judul))}{tot-1}"

    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.sheet_properties.pageSetUpPr.fitToPage = True
        w.page_setup.fitToHeight = 0
    wb.save(path)
