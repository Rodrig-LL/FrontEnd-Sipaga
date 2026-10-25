"""
Potongan satuan di luar DPP: simpanan dan iuran tetap, cicilan koperasi (Ditajenad dan LURJA),
cicilan bank, arisan Persit, simpanan pokok, batas pinjam, dan Daftar Potongan bulanan.

Keputusan juru bayar (16-17 Okt 2026, bagian 4i dokumen konteks; jangan diubah tanpa bertanya):
- hasil hitungan = berkas Daftar Potongan terpisah; buku pembukuan tetap dikosongkan G..AI;
- simpanan wajib dan potongan bertarif tetap: angka per orang yang dibawa dari buku bulan lalu,
  berubah hanya bila diubah di aplikasi; anggota baru mulai kosong;
- cicilan Koperasi Ditajenad dan LURJA dicatat saat cair (cicilan + tenor), dipotong sampai lunas;
  pinjaman yang sedang berjalan diimpor dari buku lalu dengan tenor kosong (diisi menyusul);
- cicilan bank (BRI A Yani, BRI AA, BJB, BWS, BTN) SEMENTARA angka buku lalu, sampai ada contoh
  daftar dari bank;
- arisan Persit diimpor dari DAFTAR ARISAN PERSIT tiap bulan;
- potongan > gaji: YBS disisakan Rp 5.000 (diatur di aturan), Koperasi Ditajenad yang menunggak
  (simpanan wajib lebih dulu, lalu cicilan), LURJA dan potongan lain dipotong penuh; tunggakan hanya
  dicatat, tagihan bulan depan tidak berubah; bila masih minus setelah koperasi nol: peringatan
  TINGGI, angka ditulis apa adanya;
- batas pinjam per pemberi (maks % gaji DPP, sisa gaji minimal, tenor maks) = peringatan saja;
- simpanan pokok dipotong sekali di bulan pertama keanggotaan, masuk kolom simpanan wajib;
- berkas Daftar Potongan: tabel kolom buku (baris sama dengan buku), TUNGGAKAN format buku, dan
  MASS DEBET siap kirim bank. Peringatan hanya di layar (sheet PERINGATAN tidak dipilih).

Proses gaji tidak pernah menulis ke database: daftar_proses() hanya membaca potret (snapshot()).
"""
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

import personel_db as pdb
from pembukuan import AMBANG_NAMA, _kotak, baca_buku, digit, mirip, muat_template, pastikan_xlsx

# kode, label, kelompok, judul kolom di buku (beberapa ejaan yang terukur di buku asli 16-17 Okt 2026)
JENIS = (
    ("SIMP_WAJIB", "Simpanan wajib koperasi", "koperasi", ("SIMPANAN WAJIB KOPERASI", "POT WAJIB KOPERASI")),
    ("KOPERASI", "Tanggungan Koperasi Ditajenad", "pinjaman", ("TANGGUNGAN KOPERASI",)),
    ("LURJA", "LURJA", "pinjaman", ("LURJA",)),
    ("BRI_AYANI", "BRI A Yani", "bank", ("BRI A YANI",)),
    ("BRI_AA", "BRI AA", "bank", ("BRI AA",)),
    ("BJB", "BJB", "bank", ("BJB",)),
    ("BWS", "BWS", "bank", ("BWS",)),
    ("BTN", "BTN", "bank", ("BTN",)),
    ("TWP", "TWP", "tetap", ("TWP",)),
    ("PERSIT_BINUM", "Persit Binum", "tetap", ("PERSIT BINUM",)),
    ("PERSIT_BINCAB", "Persit Bincab", "tetap", ("PERSIT BINCAB",)),
    ("PERSIT_PERSPRA", "Persit Perspra", "tetap", ("PERSIT PERSPRA",)),
    ("PERSIT_DIASAHPRA", "Persit Diasahpra", "tetap", ("PERSIT DIASAHPRA",)),
    ("PERSIT_BINPNS", "Persit Binpns", "tetap", ("PERSIT BINPNS",)),
    ("PERSIT_SIKHIB", "Persit Sikhib", "tetap", ("PERSIT SIKHIB",)),
    ("PERSIT_MINU", "Persit Minu", "tetap", ("PERSIT MINU",)),
    ("PERSIT_IT", "Persit IT", "tetap", ("PERSIT IT",)),
    ("KOWAD", "Kowad", "tetap", ("KOWAD",)),
    ("IURAN_PAMA", "Iuran Pama", "tetap", ("IURAN PAMA",)),
    ("IURAN_KASI", "Iuran Kasi", "tetap", ("IURAN KASI",)),
    ("SOSIAL_URDAL", "Sosial Urdal", "tetap", ("SOSIAL URDAL", "SOSILA URDAL")),
    ("PMI", "PMI", "tetap", ("PMI",)),
    ("KORPRI", "Korpri", "tetap", ("KORPRI",)),
    ("LAIN_LAIN", "Lain-lain", "tetap", ("LAIN-LAIN",)),
    ("ARISAN", "Arisan Persit", "arisan", ("ARISAN PERSIT",)),
)
LABEL = {k: l for k, l, _, _ in JENIS}
KELOMPOK = {k: g for k, _, g, _ in JENIS}
PINJAMAN = ("KOPERASI", "LURJA")                       # dicatat di tabel pinjaman, bukan potongan
KODE_POTONGAN = tuple(k for k, *_ in JENIS if k not in PINJAMAN)
PEMBERI_BATAS = ("KOPERASI", "LURJA", "BRI_AYANI", "BRI_AA", "BJB", "BWS", "BTN")
# MASS DEBET = kolom AU buku asli: =SUM(G:H,K:L,N:AI), semua potongan kecuali BRI A Yani, BRI AA,
# dan TWP (ketiganya dipotong sendiri oleh pihaknya). Terbukti 271/271 rekening TNI BRI Okt 2026.
LUAR_MASS_DEBET = {"BRI_AYANI", "BRI_AA", "TWP"}
BULAN = ("JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI", "AGUSTUS", "SEPTEMBER",
         "OKTOBER", "NOVEMBER", "DESEMBER")
RE_BULAN = re.compile(r"(" + "|".join(BULAN) + r")\W{0,3}(\d{4})")
RE_BANK = re.compile(r"(?<![A-Z])(BRI|BSI|BTN)(?![A-Z])")
SISA_YBS = 5000                                         # praktik buku Okt 2026: 19 orang YBS tepat 5.000


def _norm(s) -> str:
    return re.sub(r"[^0-9A-Z]", "", str(s or "").upper())


JUDUL = {_norm(a): k for k, _, _, alias in JENIS for a in alias}


def _rp(n) -> str:
    return "" if n in (None, "", 0) else f"{int(n):,}".replace(",", ".")


# ----------------------------------------------------------------------------- periode
def periode(teks) -> str:
    """'OKTOBER 2026' (atau 'BLN OKTOBER 2026 ...') -> '2026-10'; kosong bila tidak terbaca."""
    m = RE_BULAN.search(str(teks or "").upper())
    return f"{m.group(2)}-{BULAN.index(m.group(1)) + 1:02d}" if m else ""


def teks_periode(p) -> str:
    if not _periode_sah(p):
        return ""
    return f"{BULAN[int(p[5:7]) - 1]} {p[:4]}"


def _periode_sah(p) -> bool:
    return bool(re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", str(p or "")))


def ke_berapa(mulai, p) -> int:
    """Cicilan ke berapa pada periode p bila potongan pertama di bulan `mulai` (1 = bulan mulai)."""
    return (int(p[:4]) - int(mulai[:4])) * 12 + int(p[5:7]) - int(mulai[5:7]) + 1


def _ini() -> str:
    return datetime.now().strftime("%Y-%m")


# ----------------------------------------------------------------------------- isian angka
def _angka(v, nama, galat, bulat=True, boleh_kosong=True):
    """Isian rupiah/angka dari form: '5.000', 'Rp 5.000', 5000 -> 5000. Kosong -> None."""
    if v is None or (isinstance(v, str) and not v.strip()):
        if not boleh_kosong:
            galat[nama] = "Harus diisi."
        return None
    if isinstance(v, bool):
        galat[nama] = "Harus berupa angka."
        return None
    if isinstance(v, (int, float)):
        x = v
    else:
        s = re.sub(r"(?i)^\s*rp\.?\s*", "", str(v)).strip().replace(" ", "")
        s = s.replace(".", "").replace(",", ".") if bulat else s.replace(",", ".")
        try:
            x = float(s)
        except ValueError:
            galat[nama] = "Harus berupa angka."
            return None
    if x < 0:
        galat[nama] = "Tidak boleh negatif."
        return None
    if bulat:
        if x != int(x):
            galat[nama] = "Harus bilangan bulat."
            return None
        return int(x)
    return round(x, 2)


# ----------------------------------------------------------------------------- aturan
def _aturan_bawaan():
    return {"sisa_ybs": SISA_YBS, "pokok": None,
            "batas": {p: {"persen": None, "sisa_min": None, "tenor_maks": None} for p in PEMBERI_BATAS},
            "arisan": {}}


def _aturan(con) -> dict:
    a = _aturan_bawaan()
    for r in con.execute("SELECT kunci, nilai FROM aturan"):
        v = json.loads(r["nilai"])
        if r["kunci"] == "batas":
            for p in PEMBERI_BATAS:
                a["batas"][p].update((v or {}).get(p) or {})
        else:
            a[r["kunci"]] = v
    return a


def aturan() -> dict:
    with pdb._sambung() as con:
        return _aturan(con)


def _set_aturan(con, kunci, nilai):
    con.execute("INSERT INTO aturan (kunci, nilai) VALUES (?, ?) "
                "ON CONFLICT(kunci) DO UPDATE SET nilai = excluded.nilai", (kunci, json.dumps(nilai, ensure_ascii=False)))


def simpan_aturan(data: dict, operator: str) -> dict:
    """Sisa YBS, nominal simpanan pokok, dan batas pinjam per pemberi. Galat per isian."""
    g, baru = {}, {}
    baru["sisa_ybs"] = _angka(data.get("sisa_ybs"), "sisa_ybs", g, boleh_kosong=False)
    baru["pokok"] = _angka(data.get("pokok"), "pokok", g)
    if baru["pokok"] == 0:
        baru["pokok"] = None
    baru["batas"] = {}
    for p in PEMBERI_BATAS:
        b = (data.get("batas") or {}).get(p) or {}
        persen = _angka(b.get("persen"), f"batas.{p}.persen", g, bulat=False)
        if persen is not None and not 0 < persen <= 100:
            g[f"batas.{p}.persen"] = "Antara 0 dan 100."
        tenor = _angka(b.get("tenor_maks"), f"batas.{p}.tenor_maks", g)
        if tenor == 0:
            g[f"batas.{p}.tenor_maks"] = "Minimal 1 bulan, atau kosongkan."
        baru["batas"][p] = {"persen": persen, "sisa_min": _angka(b.get("sisa_min"), f"batas.{p}.sisa_min", g),
                            "tenor_maks": tenor}
    if g:
        return {"ok": False, "galat": g}
    with pdb._tulis, pdb._sambung() as con:
        lama = _aturan(con)
        beda = {}
        if lama["sisa_ybs"] != baru["sisa_ybs"]:
            beda["Sisa gaji minimal (YBS)"] = [_rp(lama["sisa_ybs"]) or "0", _rp(baru["sisa_ybs"]) or "0"]
        if lama["pokok"] != baru["pokok"]:
            beda["Simpanan pokok"] = [_rp(lama["pokok"]), _rp(baru["pokok"])]
        for p in PEMBERI_BATAS:
            for k, lbl in (("persen", "maks % gaji"), ("sisa_min", "sisa gaji minimal"), ("tenor_maks", "tenor maks")):
                a, b = lama["batas"][p][k], baru["batas"][p][k]
                if a != b:
                    f = (lambda x: "" if x is None else (f"{x:g}%" if k == "persen" else
                                                         _rp(x) if k == "sisa_min" else f"{x} bln"))
                    beda[f"Batas {LABEL[p]}: {lbl}"] = [f(a), f(b)]
        if not beda:
            return {"ok": True, "berubah": False}
        for k in ("sisa_ybs", "pokok", "batas"):
            _set_aturan(con, k, baru[k])
        pdb._catat(con, "aturan potongan", operator, None, "", "", beda)
        pdb._naik(con)
    return {"ok": True, "berubah": True}


# ----------------------------------------------------------------------------- baca
def _potongan(con, pid=None) -> dict:
    sql, arg = "SELECT personel_id, kode, nominal FROM potongan", ()
    if pid is not None:
        sql, arg = sql + " WHERE personel_id = ?", (pid,)
    hasil = {}
    for r in con.execute(sql, arg):
        hasil.setdefault(r["personel_id"], {})[r["kode"]] = r["nominal"]
    return hasil


def _pinjaman(con, pid=None) -> list:
    sql, arg = "SELECT * FROM pinjaman", ()
    if pid is not None:
        sql, arg = sql + " WHERE personel_id = ?", (pid,)
    return [dict(r) for r in con.execute(sql + " ORDER BY id", arg)]


def _pokok(con) -> dict:
    return {r["personel_id"]: dict(r) for r in con.execute("SELECT * FROM pokok")}


def jalan_pada(pj, p) -> dict:
    """Keadaan satu pinjaman pada periode p: {"dipotong": bool, "ke": n|None, "akhir": bool, "ket": teks}."""
    if pj["status"] != "aktif":
        return {"dipotong": False, "ke": None, "akhir": False, "ket": pj["status"]}
    if not _periode_sah(p):
        return {"dipotong": True, "ke": None, "akhir": False, "ket": "bulan proses tidak terbaca"}
    ke = ke_berapa(pj["mulai"], p)
    if ke < 1:
        return {"dipotong": False, "ke": ke, "akhir": False, "ket": f"mulai {teks_periode(pj['mulai']).lower()}"}
    if pj["tenor"] is None:
        return {"dipotong": True, "ke": ke, "akhir": False, "ket": "tenor belum diisi"}
    if ke > pj["tenor"]:
        return {"dipotong": False, "ke": ke, "akhir": False, "ket": "selesai"}
    return {"dipotong": True, "ke": ke, "akhir": ke == pj["tenor"], "ket": f"cicilan ke-{ke} dari {pj['tenor']}"}


def snapshot() -> dict:
    """Potret seluruh data potongan untuk proses gaji (yang tidak pernah menulis ke database)."""
    with pdb._sambung() as con:
        return {"personel": [dict(r) for r in con.execute("SELECT * FROM personel")],
                "potongan": _potongan(con), "pinjaman": _pinjaman(con), "pokok": _pokok(con),
                "aturan": _aturan(con)}


def ringkas() -> dict:
    """Untuk halaman Potongan: jenis, aturan, potongan per orang, dan semua pinjaman."""
    with pdb._sambung() as con:
        orang = {r["id"]: dict(r) for r in con.execute("SELECT id, nama, nrp_asli, pangkat, jenis, bank, status "
                                                       "FROM personel")}
        pj = _pinjaman(con)
        p = _ini()
        for x in pj:
            o = orang.get(x["personel_id"]) or {}
            x.update(nama=o.get("nama", "(dihapus)"), nrp=o.get("nrp_asli", ""), pangkat=o.get("pangkat", ""),
                     bank=o.get("bank", ""), keadaan=jalan_pada(x, p))
        pot = _potongan(con)
        hitung = {k: sum(1 for v in pot.values() if v.get(k)) for k in KODE_POTONGAN}
        return {"jenis": [{"kode": k, "label": l, "kelompok": g} for k, l, g, _ in JENIS],
                "aturan": _aturan(con), "pinjaman": pj, "potongan": pot, "hitung": hitung,
                "pokok": _pokok(con), "periode": p, "versi": pdb._versi(con)}


def data_orang(pid: int, gaji=None, p=None) -> dict:
    with pdb._sambung() as con:
        o = con.execute("SELECT * FROM personel WHERE id = ?", (pid,)).fetchone()
        if not o:
            return {"ok": False, "galat": {"_": "Personel ini sudah tidak ada."}}
        p = p if _periode_sah(p) else _ini()
        pj = _pinjaman(con, pid)
        for x in pj:
            x["keadaan"] = jalan_pada(x, p)
        pk = con.execute("SELECT * FROM pokok WHERE personel_id = ?", (pid,)).fetchone()
        a = _aturan(con)
        pot = _potongan(con, pid).get(pid, {})
    hasil = {"ok": True, "personel": dict(o), "potongan": pot, "pinjaman": pj, "pokok": dict(pk) if pk else None,
             "periode": p, "aturan": a, "gaji": gaji}
    if gaji is not None:
        hasil["batas"] = ruang_pinjam(gaji, pot, pj, (dict(pk) if pk else None), a, p)
    return hasil


# ----------------------------------------------------------------------------- batas pinjam
def _tagihan(pot, pinjaman, pokok, p) -> dict:
    """Tagihan penuh satu orang pada periode p, per kode (sebelum aturan kurang gaji)."""
    t = {k: v for k, v in pot.items() if k in KODE_POTONGAN and v}
    for pj in pinjaman:
        if jalan_pada(pj, p)["dipotong"]:
            t[pj["pemberi"]] = t.get(pj["pemberi"], 0) + pj["cicilan"]
    if pokok and pokok["bulan"] == p:
        t["SIMP_WAJIB"] = t.get("SIMP_WAJIB", 0) + pokok["nominal"]
    return t


def _cek_batas(gaji, tagihan, pinjaman, a, p) -> list:
    """Pelanggaran batas pinjam (peringatan, tidak pernah menolak)."""
    out, total = [], sum(tagihan.values())
    for pb in PEMBERI_BATAS:
        cic = tagihan.get(pb, 0)
        if not cic:
            continue
        b = a["batas"][pb]
        if b["persen"] is not None and cic > gaji * b["persen"] / 100:
            out.append(f"{LABEL[pb]}: cicilan Rp {_rp(cic)} melebihi {b['persen']:g}% gaji "
                       f"(maks Rp {_rp(int(gaji * b['persen'] / 100))}).")
        if b["sisa_min"] is not None and gaji - total < b["sisa_min"]:
            out.append(f"{LABEL[pb]}: sisa gaji sesudah semua potongan Rp {_rp(gaji - total) or '0'} di bawah "
                       f"batas minimal Rp {_rp(b['sisa_min'])}.")
        if b["tenor_maks"] is not None and pb in PINJAMAN:
            for pj in pinjaman:
                if pj["pemberi"] == pb and pj["tenor"] and pj["tenor"] > b["tenor_maks"] \
                        and jalan_pada(pj, p)["dipotong"]:
                    out.append(f"{LABEL[pb]}: tenor {pj['tenor']} bulan melebihi batas {b['tenor_maks']} bulan.")
    return out


def ruang_pinjam(gaji, pot, pinjaman, pokok, a, p) -> dict:
    """Berapa cicilan baru yang masih muat di tiap pemberi menurut batasnya (otomasi batas pinjam)."""
    t = _tagihan(pot, pinjaman, pokok, p)
    total = sum(t.values())
    per = {}
    for pb in PEMBERI_BATAS:
        b, cic, sisi = a["batas"][pb], t.get(pb, 0), []
        if b["persen"] is not None:
            sisi.append(int(gaji * b["persen"] / 100) - cic)
        if b["sisa_min"] is not None:
            sisi.append(gaji - total - b["sisa_min"])
        per[pb] = {"cicilan": cic, "ruang": max(0, min(sisi)) if sisi else None,
                   "diatur": bool(sisi) or b["tenor_maks"] is not None, "tenor_maks": b["tenor_maks"]}
    return {"gaji": gaji, "total": total, "sisa": gaji - total, "per": per,
            "peringatan": _cek_batas(gaji, t, pinjaman, a, p)}


# ----------------------------------------------------------------------------- tulis per orang
def _ada_orang(con, pid):
    return con.execute("SELECT id, nrp, nama FROM personel WHERE id = ?", (pid,)).fetchone()


def simpan_orang(pid: int, nilai: dict, operator: str) -> dict:
    """Ubah potongan tetap satu orang. Kosong/0 = tidak dipotong. Dicatat per jenis di riwayat."""
    g, baru = {}, {}
    for k in KODE_POTONGAN:
        if k in (nilai or {}):
            baru[k] = _angka(nilai[k], k, g) or 0
    asing = [k for k in (nilai or {}) if k not in KODE_POTONGAN]
    if asing:
        g["_"] = "Jenis potongan tidak dikenal: " + ", ".join(asing)
    if g:
        return {"ok": False, "galat": g}
    with pdb._tulis, pdb._sambung() as con:
        o = _ada_orang(con, pid)
        if not o:
            return {"ok": False, "galat": {"_": "Personel ini sudah tidak ada."}}
        lama = _potongan(con, pid).get(pid, {})
        beda = {LABEL[k]: [_rp(lama.get(k)), _rp(v)] for k, v in baru.items() if (lama.get(k) or 0) != v}
        if not beda:
            return {"ok": True, "berubah": False}
        for k, v in baru.items():
            _set_pot(con, pid, k, v)
        pdb._catat(con, "potongan: ubah", operator, pid, o["nrp"], o["nama"], beda)
        pdb._naik(con)
    return {"ok": True, "berubah": True}


def _set_pot(con, pid, kode, nominal):
    if nominal:
        con.execute("INSERT INTO potongan (personel_id, kode, nominal) VALUES (?, ?, ?) "
                    "ON CONFLICT(personel_id, kode) DO UPDATE SET nominal = excluded.nominal", (pid, kode, nominal))
    else:
        con.execute("DELETE FROM potongan WHERE personel_id = ? AND kode = ?", (pid, kode))


def _tgl(v):
    s = str(v or "").strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d").strftime("%Y-%m-%d") if s else ""
    except ValueError:
        return None


PJ_LABEL = {"pemberi": "Pemberi", "cicilan": "Cicilan/bulan", "mulai": "Mulai dipotong", "tenor": "Tenor",
            "tgl_cair": "Tanggal cair", "catatan": "Catatan", "status": "Status"}


def _pj_teks(k, v):
    if v in (None, ""):
        return "" if k != "tenor" else "(belum diisi)"
    return {"pemberi": lambda x: LABEL.get(x, x), "cicilan": _rp, "mulai": lambda x: teks_periode(x).lower(),
            "tenor": lambda x: f"{x} bulan"}.get(k, str)(v)


def simpan_pinjaman(data: dict, operator: str, pid_pj=None, gaji=None, p=None) -> dict:
    """Tambah (pid_pj None) atau ubah satu pinjaman koperasi/LURJA. Lewat batas = peringatan saja."""
    g = {}
    try:
        pid = int(data.get("personel_id") or 0)
    except (TypeError, ValueError):
        pid = 0
    pemberi = str(data.get("pemberi") or "")
    if pemberi not in PINJAMAN:
        g["pemberi"] = "Pilih Koperasi Ditajenad atau LURJA."
    cicilan = _angka(data.get("cicilan"), "cicilan", g, boleh_kosong=False)
    if cicilan == 0:
        g["cicilan"] = "Cicilan harus lebih dari 0."
    mulai = str(data.get("mulai") or "").strip()[:7]
    if not _periode_sah(mulai):
        g["mulai"] = "Pilih bulan potongan pertama."
    tenor = _angka(data.get("tenor"), "tenor", g)
    if tenor == 0:
        g["tenor"] = "Minimal 1 bulan, atau kosongkan bila belum diketahui."
    tgl_cair = _tgl(data.get("tgl_cair"))
    if tgl_cair is None:
        g["tgl_cair"] = "Tanggal tidak sah."
    status = str(data.get("status") or "aktif")
    if status not in ("aktif", "lunas", "batal"):
        g["status"] = "Status harus aktif, lunas, atau batal."
    if g:
        return {"ok": False, "galat": g}
    v = {"pemberi": pemberi, "cicilan": cicilan, "mulai": mulai, "tenor": tenor, "tgl_cair": tgl_cair,
         "catatan": str(data.get("catatan") or "").strip(), "status": status}
    with pdb._tulis, pdb._sambung() as con:
        if pid_pj is not None:
            lama = con.execute("SELECT * FROM pinjaman WHERE id = ?", (pid_pj,)).fetchone()
            if not lama:
                return {"ok": False, "galat": {"_": "Pinjaman ini sudah tidak ada."}}
            pid = lama["personel_id"]
        o = _ada_orang(con, pid)
        if not o:
            return {"ok": False, "galat": {"_": "Personel ini sudah tidak ada."}}
        now = pdb._sekarang()
        if pid_pj is None:
            cur = con.execute("INSERT INTO pinjaman (personel_id, pemberi, cicilan, mulai, tenor, tgl_cair, catatan, "
                              "status, dibuat, diubah) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                              (pid, *[v[k] for k in ("pemberi", "cicilan", "mulai", "tenor", "tgl_cair", "catatan",
                                                     "status")], now, now))
            pid_pj = cur.lastrowid
            aksi, beda = "pinjaman: tambah", {PJ_LABEL[k]: ["", _pj_teks(k, x)] for k, x in v.items()
                                              if x not in (None, "") and k != "status"}
        else:
            beda = {PJ_LABEL[k]: [_pj_teks(k, lama[k]), _pj_teks(k, x)] for k, x in v.items() if lama[k] != x}
            if not beda:
                return {"ok": True, "id": pid_pj, "berubah": False, "peringatan": []}
            con.execute(f"UPDATE pinjaman SET {', '.join(k + ' = ?' for k in v)}, diubah = ? WHERE id = ?",
                        [*v.values(), now, pid_pj])
            aksi = (f"pinjaman: {status}" if lama["status"] != status else "pinjaman: ubah")
        pdb._catat(con, aksi, operator, pid, o["nrp"], o["nama"], beda)
        pdb._naik(con)
        peringatan = []
        if gaji is not None:
            p = p if _periode_sah(p) else _ini()
            peringatan = _cek_batas(gaji, _tagihan(_potongan(con, pid).get(pid, {}), _pinjaman(con, pid),
                                                   None, p), _pinjaman(con, pid), _aturan(con), p)
    return {"ok": True, "id": pid_pj, "berubah": True, "peringatan": peringatan}


def simpan_pokok(pid: int, bulan, nominal, operator: str) -> dict:
    """Anggota koperasi baru: simpanan pokok dipotong sekali di `bulan`. Bulan kosong = dibatalkan."""
    g = {}
    bulan = str(bulan or "").strip()[:7]
    if bulan and not _periode_sah(bulan):
        g["bulan"] = "Pilih bulan."
    n = _angka(nominal, "nominal", g)
    with pdb._tulis, pdb._sambung() as con:
        if bulan and not n:
            n = _aturan(con)["pokok"]
            if not n:
                g["nominal"] = "Isi nominal simpanan pokok (atau atur nominal bawaannya di Potongan > Aturan)."
        if g:
            return {"ok": False, "galat": g}
        o = _ada_orang(con, pid)
        if not o:
            return {"ok": False, "galat": {"_": "Personel ini sudah tidak ada."}}
        lama = con.execute("SELECT * FROM pokok WHERE personel_id = ?", (pid,)).fetchone()
        teks = lambda r: f"{_rp(r['nominal'])} ({teks_periode(r['bulan']).lower()})" if r else ""  # noqa: E731
        baru = {"bulan": bulan, "nominal": n} if bulan else None
        if (teks(lama) if lama else "") == (teks(baru) if baru else ""):
            return {"ok": True, "berubah": False}
        if baru:
            con.execute("INSERT INTO pokok (personel_id, bulan, nominal) VALUES (?, ?, ?) ON CONFLICT(personel_id) "
                        "DO UPDATE SET bulan = excluded.bulan, nominal = excluded.nominal", (pid, bulan, n))
        else:
            con.execute("DELETE FROM pokok WHERE personel_id = ?", (pid,))
        pdb._catat(con, "potongan: simpanan pokok", operator, pid, o["nrp"], o["nama"],
                   {"Simpanan pokok": [teks(lama) if lama else "", teks(baru) if baru else ""]})
        pdb._naik(con)
    return {"ok": True, "berubah": True}


# ----------------------------------------------------------------------------- potongan massal
def _saring(personel, potongan, s: dict) -> list:
    hasil = []
    pangkat = {_norm(x) for x in (s.get("pangkat") or []) if _norm(x)}
    for p in personel:
        if (s.get("status") or "aktif") == "aktif" and p["status"] != "aktif":
            continue
        if s.get("jenis") and p["jenis"] != s["jenis"]:
            continue
        if s.get("bank") and (p["bank"] or "-") != s["bank"]:
            continue
        if pangkat and _norm((p["pangkat"].split() or [""])[0]) not in pangkat:
            continue
        punya = bool(potongan.get(p["id"], {}).get(s.get("kode")))
        if s.get("punya") == "ya" and not punya or s.get("punya") == "tidak" and punya:
            continue
        hasil.append(p)
    return hasil


def rencana_massal(data: dict) -> dict:
    """Pratinjau satu jenis potongan untuk sekelompok personel. Nominal 0/kosong = dihapus."""
    g = {}
    kode = str(data.get("kode") or "")
    if kode not in KODE_POTONGAN:
        g["kode"] = "Pilih jenis potongan."
    nominal = _angka(data.get("nominal"), "nominal", g) or 0
    if g:
        return {"ok": False, "galat": g}
    s = {**(data.get("saring") or {}), "kode": kode}
    with pdb._sambung() as con:
        versi = pdb._versi(con)
        personel = [dict(r) for r in con.execute("SELECT * FROM personel ORDER BY jenis DESC, nama")]
        pot = _potongan(con)
    ubah, tetap = [], 0
    for p in _saring(personel, pot, s):
        lama = pot.get(p["id"], {}).get(kode, 0)
        if lama == nominal:
            tetap += 1
        else:
            ubah.append({"id": p["id"], "nama": p["nama"], "nrp": p["nrp_asli"], "pangkat": p["pangkat"],
                         "lama": lama, "baru": nominal})
    return {"ok": True, "rencana": {"jenis_impor": "massal", "versi": versi, "kode": kode, "label": LABEL[kode],
                                    "nominal": nominal, "saring": s, "ubah": ubah, "tetap": tetap}}


# ----------------------------------------------------------------------------- impor
def _pasangkan(baris, calon):
    """Pasangkan baris (buku/daftar/DPP) dengan personel: NRP dulu, lalu nama (>= AMBANG_NAMA dan yang
    terbaik harus unik). Mengembalikan list (baris, personel|None, cara)."""
    per_nrp = {}
    for p in calon:
        per_nrp.setdefault(p["nrp"], p)
    out = []
    for b in baris:
        p, cara = per_nrp.get(digit(b["nrp"])), "NRP"
        if not p and b.get("nama"):
            skor = sorted(((mirip(b["nama"], q["nama"]), q["id"], q) for q in calon), key=lambda x: (-x[0], x[1]))
            if skor and skor[0][0] >= AMBANG_NAMA and (len(skor) == 1 or skor[1][0] < skor[0][0]):
                p, cara = skor[0][2], "nama"
        out.append((b, p, cara if p else ""))
    return out


def _rows(path, sheet, sampai):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb[sheet].iter_rows(min_row=1, max_row=sampai, values_only=True))
    finally:
        wb.close()


def kolom_buku(rows, header: int) -> dict:
    """Kolom potongan satu sheet buku dari judul di baris header: {"peta": {huruf: kode}, "judul": {huruf:
    teks}, "asing": {huruf: teks}, "akhir": huruf kolom terakhir sebelum JUMLAH, "ybs": huruf YBS}."""
    judul, peta, asing, akhir, ybs = {}, {}, {}, None, None
    baris = rows[header - 1] if 0 < header <= len(rows) else ()
    for j in range(6, len(baris)):                     # mulai G
        t = str(baris[j] or "").strip()
        h = get_column_letter(j + 1)
        n = _norm(t)
        if n == "JUMLAH" and akhir is None:
            akhir = get_column_letter(j)
            continue
        if n == "YBS" and ybs is None:
            ybs = h
        if akhir is not None:
            continue
        judul[h] = t
        if n in JUDUL:
            peta[h] = JUDUL[n]
        elif t:
            asing[h] = t
    return {"peta": peta, "judul": judul, "asing": asing, "akhir": akhir or "AI", "ybs": ybs}


def _bilangan(v):
    if isinstance(v, bool) or v is None:
        return 0
    if isinstance(v, (int, float)):
        return int(round(v))
    s = digit(v)
    return int(s) if s and str(v).strip().replace(".", "").replace(",", "").isdigit() else 0


def _sheet_buku(sumber, tpl, sheet, jenis):
    """(t, buku, rows) satu sheet pembukuan: BUKU TNI dengan peta template, BUKU PNS dengan kolom
    rekening dari judulnya dan JUMLAH di kolom B (sama dengan personel_db.rencana_bank)."""
    rows25 = _rows(sumber, sheet, 25)
    t = pdb._tpl_sheet(tpl, sheet, rows25)
    if t is None:
        return None
    if jenis == "PNS":
        t["akhir_tabel"] = {**t["akhir_tabel"], "kolom": "B"}
    buku = baca_buku(sumber, t, boleh_kosong=True)
    return t, buku, _rows(sumber, sheet, buku["baris_jumlah"])


def rencana_buku(paths) -> dict:
    """Pratinjau impor potongan dari pembukuan bulan lalu (BUKU TNI dan BUKU PNS): potongan tetap dan
    cicilan bank per orang, plus pinjaman koperasi/LURJA yang sedang berjalan (tenor belum diisi)."""
    tpl = muat_template()
    with pdb._sambung() as con:
        versi = pdb._versi(con)
        semua = [dict(r) for r in con.execute("SELECT * FROM personel")]
        pot = _potongan(con)
        pj = _pinjaman(con)
        a = _aturan(con)
    hasil = {"jenis_impor": "potongan", "versi": versi, "berkas": [], "ubah": [], "tetap": 0, "pinjaman_baru": [],
             "pinjaman_beda": [], "tidak_ketemu": [], "masalah": [], "catatan": []}
    usul = {}
    for path in paths:
        path = Path(path)
        sumber = pastikan_xlsx(path)
        wb = load_workbook(sumber, read_only=True)
        try:
            ada = [s for s in ("BUKU TNI", "BUKU PNS") if s in wb.sheetnames]
        finally:
            wb.close()
        if not ada:
            raise SystemExit(_kotak("bukan berkas pembukuan", [f"Berkas : {path.name}",
                                                               "Tidak ada sheet BUKU TNI maupun BUKU PNS."],
                                    "Pilih berkas PEMBUKUAN GAJI bulan lalu."))
        info = {"file": path.name, "sheet": []}
        for sheet in ada:
            jenis = "PNS" if sheet == "BUKU PNS" else "TNI"
            sb = _sheet_buku(sumber, tpl, sheet, jenis)
            if sb is None:
                hasil["catatan"].append(f"{path.name}, sheet {sheet}: kolom NO REKENING tidak ditemukan; dilewati.")
                continue
            t, buku, rows = sb
            kb = kolom_buku(rows, buku["header"])
            p_buku = periode(buku["bulan"]) or _ini()
            if not periode(buku["bulan"]):
                hasil["catatan"].append(f"{path.name}, {sheet}: bulan buku tidak terbaca; pinjaman berjalan dianggap "
                                        f"mulai {teks_periode(p_buku).lower()}.")
            info["sheet"].append({"sheet": sheet, "baris": len(buku["anggota"]), "bulan": buku["bulan"] or "",
                                  "kolom": len(kb["peta"])})
            if kb["asing"]:
                isi = [h for h in kb["asing"] if any(_bilangan(rows[b["baris"] - 1][column_index_from_string(h) - 1])
                                                     for b in buku["anggota"])]
                if isi:
                    hasil["catatan"].append(f"{path.name}, {sheet}: kolom " + ", ".join(
                        f"{h} ({kb['asing'][h]!r})" for h in isi) + " berisi angka tetapi judulnya tidak dikenal, "
                        "jadi tidak diimpor. Isi lewat form potongan bila perlu.")
            if "ARISAN" in kb["peta"].values():
                hasil["catatan"].append(f"{path.name}, {sheet}: kolom ARISAN PERSIT tidak diimpor dari buku; arisan "
                                        "diambil dari DAFTAR ARISAN PERSIT tiap bulan.")
            calon = [p for p in semua if p["jenis"] == jenis]
            for b, p, cara in _pasangkan(buku["anggota"], calon):
                if not p:
                    hasil["tidak_ketemu"].append({"file": path.name, "sheet": sheet, "baris": b["baris"],
                                                  "nama": b["nama"], "nrp": b["nrp_asli"]})
                    continue
                r = rows[b["baris"] - 1]
                nilai = {}
                for h, kode in kb["peta"].items():
                    j = column_index_from_string(h) - 1
                    nilai[kode] = nilai.get(kode, 0) + (_bilangan(r[j]) if j < len(r) else 0)
                ybs = None
                if kb["ybs"]:
                    j = column_index_from_string(kb["ybs"]) - 1
                    ybs = r[j] if j < len(r) and isinstance(r[j], (int, float)) else None
                usul.setdefault(p["id"], []).append({"nilai": nilai, "ybs": ybs, "periode": p_buku, "cara": cara,
                                                     "asal": f"{path.name} {sheet} baris {b['baris']}"})
        hasil["berkas"].append(info)
    per_id = {p["id"]: p for p in semua}
    for pid, us in usul.items():
        p = per_id[pid]
        if len(us) > 1:
            hasil["masalah"].append({"id": pid, "nama": p["nama"], "nrp": p["nrp_asli"],
                                     "alasan": "Tercatat di lebih dari satu baris pembukuan: "
                                               + "; ".join(u["asal"] for u in us) + ". Potongannya tidak diimpor; "
                                               "isi lewat form potongan."})
            continue
        u = us[0]
        lama = pot.get(pid, {})
        baru = {k: v for k, v in u["nilai"].items() if k in KODE_POTONGAN and k != "ARISAN"}
        beda = {k: [lama.get(k, 0), v] for k, v in baru.items() if lama.get(k, 0) != v}
        if beda:
            hasil["ubah"].append({"id": pid, "nama": p["nama"], "nrp": p["nrp_asli"], "asal": u["asal"],
                                  "cara": u["cara"], "potongan": {k: b for k, (_, b) in beda.items()},
                                  "perubahan": {LABEL[k]: [_rp(x), _rp(y)] for k, (x, y) in beda.items()}})
        else:
            hasil["tetap"] += 1
        for pb in PINJAMAN:
            cic = u["nilai"].get(pb, 0)
            aktif = [x for x in pj if x["personel_id"] == pid and x["pemberi"] == pb and x["status"] == "aktif"]
            if aktif:
                jalan = sum(x["cicilan"] for x in aktif if jalan_pada(x, u["periode"])["dipotong"])
                if jalan != cic:
                    hasil["pinjaman_beda"].append({"id": pid, "nama": p["nama"], "nrp": p["nrp_asli"],
                                                   "pemberi": pb, "buku": cic, "database": jalan})
                continue
            if not cic:
                continue
            tanda = ""
            if pb == "KOPERASI" and u["ybs"] is not None and u["ybs"] <= a["sisa_ybs"]:
                tanda = (f"Bulan lalu gaji tidak cukup (YBS Rp {_rp(u['ybs']) or '0'}), jadi angka di buku mungkin "
                         "sudah dikurangi; cicilan sebenarnya bisa lebih besar. Periksa di sheet TUNGGAKAN.")
            hasil["pinjaman_baru"].append({"id": pid, "nama": p["nama"], "nrp": p["nrp_asli"], "pemberi": pb,
                                           "cicilan": cic, "mulai": u["periode"], "asal": u["asal"], "tandai": tanda})
    return hasil


def _judul_arisan(rows) -> str:
    for r in rows[:5]:
        for v in r:
            if "ARISAN" in str(v or "").upper():
                return str(v).strip()
    return ""


def rencana_arisan(paths) -> dict:
    """Pratinjau impor DAFTAR ARISAN PERSIT bulan ini (sheet TNI dan PNS). Bank dari judul ('... GAJI DI
    BRI'), bila tidak ada dari nama berkas. Yang tidak tercantum lagi di daftar bank itu berhenti arisan."""
    with pdb._sambung() as con:
        versi = pdb._versi(con)
        semua = [dict(r) for r in con.execute("SELECT * FROM personel")]
        pot = _potongan(con)
    hasil = {"jenis_impor": "arisan", "versi": versi, "berkas": [], "ubah": [], "berhenti": [], "tetap": 0,
             "tidak_ketemu": [], "masalah": [], "catatan": [], "meta": {}}
    daftar = []
    for path in paths:
        path = Path(path)
        sumber = pastikan_xlsx(path)
        wb = load_workbook(sumber, read_only=True, data_only=True)
        try:
            sheets = {ws.title: list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 2000, 2000), max_col=12,
                                                  values_only=True)) for ws in wb.worksheets}
        finally:
            wb.close()
        judul, bank, info = "", None, {"file": path.name, "sheet": []}
        for nama, rows in sheets.items():
            hdr = None
            for i, r in enumerate(rows[:15]):
                n = [_norm(v) for v in r]
                kol_nrp = next((j for j, x in enumerate(n) if x in ("NRPNIP", "NRP", "NIP")), None)
                kol_pot = next((j for j, x in enumerate(n) if re.fullmatch(r"P[O0]TONGAN", x or "")), None)
                if kol_nrp is not None and kol_pot is not None:
                    hdr = (i, kol_nrp, kol_pot, next((j for j, x in enumerate(n) if x == "NAMA"), 1))
                    break
            if hdr is None:
                continue
            j_sheet = _judul_arisan(rows)
            judul = judul or j_sheet
            m = RE_BANK.search(j_sheet.upper())
            if m:
                bank = bank or m.group(1)
            n_baris = 0
            for i in range(hdr[0] + 1, len(rows)):
                r = rows[i]
                nrp = digit(r[hdr[1]] if hdr[1] < len(r) else "")
                nominal = _bilangan(r[hdr[2]] if hdr[2] < len(r) else 0)
                if len(nrp) < pdb.MIN_DIGIT or not nominal:
                    continue
                n_baris += 1
                daftar.append({"nrp": nrp, "nrp_asli": str(r[hdr[1]]).strip(), "nama": str(r[hdr[3]] or "").strip(),
                               "nominal": nominal, "asal": f"{path.name} sheet {nama} baris {i + 1}"})
            info["sheet"].append({"sheet": nama, "baris": n_baris})
        if not info["sheet"]:
            raise SystemExit(_kotak("bukan daftar arisan", [f"Berkas : {path.name}",
                                                            "Tidak ada sheet dengan judul kolom NRP/NIP dan POTONGAN."],
                                    "Pilih berkas DAFTAR ARISAN PERSIT bulan ini."))
        if not bank:
            m = RE_BANK.search(path.stem.upper())
            bank = m.group(1) if m else None
        if not bank:
            raise SystemExit(_kotak("bank daftar arisan tidak diketahui",
                                    [f"Berkas : {path.name}", f"Judul  : {judul or '(tidak ada)'}",
                                     "Judul maupun nama berkasnya tidak menyebut BRI, BSI, atau BTN. Bank perlu",
                                     "diketahui supaya arisan personel bank lain tidak ikut berhenti."],
                                    "Ganti nama berkasnya, mis. DAFTAR ARISAN PERSIT BRI.xls, lalu coba lagi."))
        info.update(judul=judul, bank=bank)
        hasil["berkas"].append(info)
        hasil["meta"][bank] = {"file": path.name, "judul": judul, "jumlah": sum(s["baris"] for s in info["sheet"])}
        for d in daftar:
            d.setdefault("bank", bank)
    bank_daftar = set(hasil["meta"])
    baru, dilihat = {}, {}
    for d, p, cara in _pasangkan(daftar, semua):
        if not p:
            hasil["tidak_ketemu"].append({"nama": d["nama"], "nrp": d["nrp_asli"], "asal": d["asal"]})
            continue
        if p["id"] in dilihat:
            hasil["masalah"].append({"id": p["id"], "nama": p["nama"], "nrp": p["nrp_asli"],
                                     "alasan": f"Tercantum dua kali ({dilihat[p['id']]}; {d['asal']}). "
                                               "Yang pertama dipakai."})
            continue
        dilihat[p["id"]] = d["asal"]
        baru[p["id"]] = d["nominal"]
        if p["bank"] and p["bank"] != d["bank"]:
            hasil["catatan"].append(f"{p['nama']} ({p['nrp_asli']}) tercantum di daftar arisan {d['bank']}, tetapi "
                                    f"bank-nya di database {p['bank']}.")
        lama = pot.get(p["id"], {}).get("ARISAN", 0)
        if lama == d["nominal"]:
            hasil["tetap"] += 1
        else:
            hasil["ubah"].append({"id": p["id"], "nama": p["nama"], "nrp": p["nrp_asli"], "lama": lama,
                                  "baru": d["nominal"], "cara": cara, "asal": d["asal"]})
    for p in semua:
        lama = pot.get(p["id"], {}).get("ARISAN", 0)
        if lama and p["id"] not in baru and p["bank"] in bank_daftar:
            hasil["berhenti"].append({"id": p["id"], "nama": p["nama"], "nrp": p["nrp_asli"], "lama": lama,
                                      "bank": p["bank"]})
    return hasil


def terapkan(rencana: dict, operator: str) -> dict:
    """Terapkan rencana impor potongan/arisan atau potongan massal. Ditolak bila database berubah."""
    j = rencana["jenis_impor"]
    with pdb._tulis, pdb._sambung() as con:
        if pdb._versi(con) != rencana["versi"]:
            return {"ok": False, "galat": {"_": "Data berubah sejak pratinjau dibuat (mungkin di tab lain). "
                                                "Buat pratinjau ulang, lalu terapkan lagi."}}
        orang = {r["id"]: r for r in con.execute("SELECT id, nrp, nama FROM personel")}
        n = {"ubah": 0, "pinjaman": 0, "berhenti": 0}
        now = pdb._sekarang()
        if j == "potongan":
            for x in rencana["ubah"]:
                if x["id"] not in orang:
                    continue
                for k, v in x["potongan"].items():
                    _set_pot(con, x["id"], k, v)
                pdb._catat(con, "impor: potongan", operator, x["id"], orang[x["id"]]["nrp"], x["nama"], x["perubahan"])
                n["ubah"] += 1
            for x in rencana["pinjaman_baru"]:
                if x["id"] not in orang:
                    continue
                cat = f"Diimpor dari {x['asal']}: pinjaman berjalan, sisa tenor belum diketahui."
                if x["tandai"]:
                    cat += " " + x["tandai"]
                con.execute("INSERT INTO pinjaman (personel_id, pemberi, cicilan, mulai, tenor, tgl_cair, catatan, "
                            "status, dibuat, diubah) VALUES (?, ?, ?, ?, NULL, '', ?, 'aktif', ?, ?)",
                            (x["id"], x["pemberi"], x["cicilan"], x["mulai"], cat, now, now))
                pdb._catat(con, "impor: pinjaman", operator, x["id"], orang[x["id"]]["nrp"], x["nama"],
                           {"Pemberi": ["", LABEL[x["pemberi"]]], "Cicilan/bulan": ["", _rp(x["cicilan"])],
                            "Tenor": ["", "(belum diisi)"]})
                n["pinjaman"] += 1
        elif j == "arisan":
            for x in rencana["ubah"]:
                if x["id"] in orang:
                    _set_pot(con, x["id"], "ARISAN", x["baru"])
                    pdb._catat(con, "impor: arisan", operator, x["id"], orang[x["id"]]["nrp"], x["nama"],
                               {"Arisan Persit": [_rp(x["lama"]), _rp(x["baru"])]})
                    n["ubah"] += 1
            for x in rencana["berhenti"]:
                if x["id"] in orang:
                    _set_pot(con, x["id"], "ARISAN", 0)
                    pdb._catat(con, "impor: arisan", operator, x["id"], orang[x["id"]]["nrp"], x["nama"],
                               {"Arisan Persit": [_rp(x["lama"]), ""]})
                    n["berhenti"] += 1
            meta = _aturan(con)["arisan"] or {}
            for bank, m in rencana["meta"].items():
                meta[bank] = {**m, "diimpor": now}
            _set_aturan(con, "arisan", meta)
        elif j == "massal":
            for x in rencana["ubah"]:
                if x["id"] in orang:
                    _set_pot(con, x["id"], rencana["kode"], x["baru"])
                    pdb._catat(con, "potongan massal", operator, x["id"], orang[x["id"]]["nrp"], x["nama"],
                               {rencana["label"]: [_rp(x["lama"]), _rp(x["baru"])]})
                    n["ubah"] += 1
        else:
            return {"ok": False, "galat": {"_": f"Jenis impor tidak dikenal: {j}"}}
        ringkas_ = {"jenis": j, **n}
        if j == "massal":
            ringkas_.update(potongan=rencana["label"], nominal=rencana["nominal"])
        else:
            ringkas_["berkas"] = [b["file"] for b in rencana["berkas"]]
        pdb._catat(con, "impor" if j != "massal" else "potongan massal", operator, None, "", "", ringkas_)
        pdb._naik(con)
    return {"ok": True, **n}


# ----------------------------------------------------------------------------- hitung bulanan
def _hitung_satu(gaji, tagihan, sisa_ybs) -> dict:
    """Aturan kurang gaji (keputusan 17 Okt 2026): semua selain Koperasi Ditajenad dipotong penuh;
    koperasi mengambil sisanya sampai YBS tinggal `sisa_ybs`, simpanan wajib lebih dulu, lalu cicilan."""
    sw_t, kp_t = tagihan.get("SIMP_WAJIB", 0), tagihan.get("KOPERASI", 0)
    lain = sum(v for k, v in tagihan.items() if k not in ("SIMP_WAJIB", "KOPERASI"))
    tersedia = max(0, gaji - lain - sisa_ybs)
    sw = min(sw_t, tersedia)
    kp = min(kp_t, tersedia - sw)
    pot = {k: v for k, v in tagihan.items() if v}
    for k, v in (("SIMP_WAJIB", sw), ("KOPERASI", kp)):
        if v:
            pot[k] = v
        else:
            pot.pop(k, None)
    total = sum(pot.values())
    return {"potongan": pot, "total": total, "ybs": gaji - total,
            "kop": {"tagihan": sw_t + kp_t, "sw": sw, "kp": kp, "terpotong": sw + kp,
                    "tunggakan": sw_t + kp_t - sw - kp},
            "mass_debet": sum(v for k, v in pot.items() if k not in LUAR_MASS_DEBET)}


def hitung_daftar(snap: dict, bulan_teks: str, bank, kelompok: list) -> dict:
    """Hitung potongan bulan ini. kelompok = [{"sheet", "jenis", "baris": [{"baris", "no", "nama",
    "pangkat", "nrp", "nrp_asli", "rekening", "nama_rekening", "gaji"}], ...}]; gaji None = tidak ada di
    DPP (potongan tidak dihitung). Tidak menulis ke database."""
    a, p = snap["aturan"], periode(bulan_teks)
    sisa_ybs = a["sisa_ybs"]
    per_jenis = {}
    for o in snap["personel"]:
        per_jenis.setdefault(o["jenis"], []).append(o)
    pj_orang = {}
    for x in snap["pinjaman"]:
        pj_orang.setdefault(x["personel_id"], []).append(x)
    peringatan, out, gaji_pid, dipakai = [], [], {}, {}

    def ingat(tingkat, kategori, detail, r=None, sheet=""):
        peringatan.append({"tingkat": tingkat, "kategori": kategori, "detail": detail,
                           "nama": r["nama"] if r else "", "nrp": r["nrp_asli"] if r else "",
                           "letak": (f"{sheet} baris {r['baris']}" if r and r.get("baris") else sheet)})

    if not p:
        ingat("TINGGI", "Bulan proses tidak terbaca",
              "Bulan DPP tidak terbaca, jadi jadwal cicilan dan simpanan pokok tidak bisa ditentukan. Semua "
              "pinjaman aktif dianggap dipotong; periksa hasilnya.")
    for kel in kelompok:
        sheet, jenis = kel["sheet"], kel["jenis"]
        baris_out = []
        for r, o, cara in _pasangkan(kel["baris"], per_jenis.get(jenis, [])):
            x = {**r, "pid": o["id"] if o else None, "cara": cara, "ket": [], "potongan": {}, "total": None,
                 "ybs": None, "kop": None, "mass_debet": 0, "pinjaman": []}
            baris_out.append(x)
            if o and o["id"] in dipakai:
                ingat("SEDANG", "Personel tercatat di dua baris",
                      f"Sama dengan {dipakai[o['id']]}. Potongannya hanya dihitung di baris pertama.", r, sheet)
                x["ket"].append("GANDA")
                o = None
                x["pid"] = None
            if not o:
                if "GANDA" not in x["ket"]:
                    ingat("SEDANG", "Tidak ada di database personel",
                          "Tidak cocok dengan siapa pun di database personel (NRP maupun nama), jadi potongannya "
                          "tidak diketahui dan dibiarkan kosong. Tambahkan lewat Data Personel.", r, sheet)
                    x["ket"].append("TIDAK ADA DI DATABASE")
                continue
            dipakai[o["id"]] = f"{sheet} baris {r['baris']}" if r.get("baris") else sheet
            x["rekening"] = o["rekening"] or r.get("rekening") or ""
            x["nama_rekening"] = o["nama_rekening"] or r.get("nama_rekening") or r["nama"]
            pinj = pj_orang.get(o["id"], [])
            pot = snap["potongan"].get(o["id"], {})
            pokok = snap["pokok"].get(o["id"])
            tagihan = _tagihan(pot, pinj, pokok, p)
            for pjx in pinj:
                k = jalan_pada(pjx, p)
                if k["dipotong"]:
                    x["pinjaman"].append({"pemberi": pjx["pemberi"], "cicilan": pjx["cicilan"], "ket": k["ket"]})
                    if pjx["tenor"] is None:
                        x["ket"].append("TENOR BELUM DIISI")
                    elif k["akhir"]:
                        ingat("INFO", "Cicilan terakhir", f"{LABEL[pjx['pemberi']]} Rp {_rp(pjx['cicilan'])}: "
                              f"{k['ket']}. Bulan depan tidak dipotong lagi.", r, sheet)
            if pokok and pokok["bulan"] == p:
                ingat("INFO", "Simpanan pokok anggota baru", f"Rp {_rp(pokok['nominal'])} ditambahkan ke simpanan "
                      "wajib bulan ini.", r, sheet)
            if r.get("gaji") is None:
                if tagihan:
                    ingat("SEDANG", "Tidak ada di DPP", f"Punya potongan Rp {_rp(sum(tagihan.values()))}, tetapi "
                          "tidak ada di DPP bulan ini, jadi gajinya tidak diketahui. Potongannya tidak dihitung "
                          "dan tidak masuk MASS DEBET.", r, sheet)
                x["ket"].append("TIDAK ADA DI DPP")
                continue
            gaji = int(r["gaji"])
            gaji_pid[o["id"]] = gaji
            h = _hitung_satu(gaji, tagihan, sisa_ybs)
            x.update(potongan=h["potongan"], total=h["total"], ybs=h["ybs"], kop=h["kop"], mass_debet=h["mass_debet"])
            if h["kop"]["tunggakan"]:
                x["ket"].append("TUNGGAKAN")
            if h["ybs"] < 0:
                x["ket"].append("YBS MINUS")
                ingat("TINGGI", "Potongan melebihi gaji",
                      f"Gaji Rp {_rp(gaji)}, potongan selain Koperasi Ditajenad saja sudah Rp {_rp(h['total'])}: "
                      f"YBS minus Rp {_rp(-h['ybs'])}. Ditulis apa adanya; putuskan sendiri potongan mana yang "
                      "ditunda.", r, sheet)
            elif h["ybs"] < sisa_ybs:
                ingat("SEDANG", "Sisa gaji di bawah batas",
                      f"YBS Rp {_rp(h['ybs']) or '0'}, di bawah Rp {_rp(sisa_ybs)}, walaupun Koperasi Ditajenad "
                      "sudah tidak dipotong.", r, sheet)
            for t in _cek_batas(gaji, tagihan, pinj, a, p):
                ingat("SEDANG", "Lewat batas pinjam", t, r, sheet)
        out.append({"sheet": sheet, "jenis": jenis, "baris": baris_out})
    tenor_kosong = [x for k in out for x in k["baris"] if "TENOR BELUM DIISI" in x["ket"]]
    if tenor_kosong:
        ingat("SEDANG", "Tenor pinjaman belum diisi",
              f"{len(tenor_kosong)} orang punya pinjaman koperasi/LURJA yang sisa tenornya belum diisi. Cicilannya "
              "tetap dipotong tiap bulan sampai tenornya diisi (Potongan > Pinjaman).")
    return {"bulan": bulan_teks, "periode": p, "bank": bank, "kelompok": out, "peringatan": peringatan,
            "gaji": gaji_pid, "aturan": {"sisa_ybs": sisa_ybs}}


# ----------------------------------------------------------------------------- tulis berkas
TEBAL = Font(bold=True)
MERAH = Font(color="9A2E17", bold=True)
KUNING = PatternFill("solid", fgColor="FFF2CC")
ABU = PatternFill("solid", fgColor="EDEDED")
RP = "#,##0"


def _lebar(ws, lebar):
    for h, w in lebar.items():
        ws.column_dimensions[h].width = w


def _sheet_potongan(wb, kel, kb, header, baris_jumlah, bank, bulan):
    ws = wb.create_sheet(f"POTONGAN {kel['jenis']}")
    akhir = column_index_from_string(kb["akhir"])
    kolom = [get_column_letter(j) for j in range(7, akhir + 1)]
    j_jml = akhir + 1
    hj, hy, ht, hk = (get_column_letter(j_jml + i) for i in range(4))
    data = [x for x in kel["baris"] if x.get("baris")]
    lain = [x for x in kel["baris"] if not x.get("baris")]
    if header >= 4:
        ws["A1"] = f"DAFTAR POTONGAN SATUAN - {kel['sheet']} - {bank or ''} - BULAN {bulan}".replace(" -  -", " -")
        ws["A1"].font = Font(bold=True, size=13)
        if data:
            a, b = min(x["baris"] for x in data), max(x["baris"] for x in data)
            ws["A2"] = (f"Salin sel {kolom[0]}{a}:{kolom[-1]}{b} lalu tempel sebagai NILAI (Paste Special > Values) "
                        f"ke sheet {kel['sheet']} pada sel yang sama. Baris dan kolomnya sama dengan buku.")
        ws["A3"] = (f"Dihitung SIPAGA {datetime.now():%d-%m-%Y %H:%M}. Baris tanpa angka: tidak ada di DPP atau "
                    "tidak ada di database personel (lihat KET).")
    judul = ["NO", "NAMA", "PANGKAT", "", "NRP/NIP", "GAJI"]
    for j, t in enumerate(judul, 1):
        ws.cell(header, j, t)
    for h in kolom:
        t = kb["judul"].get(h) or ""
        ws[f"{h}{header}"] = t or (LABEL[kb["peta"][h]].upper() if h in kb["peta"] else "")
    for h, t in ((hj, "JUMLAH POTONGAN"), (hy, "YBS"), (ht, "TUNGGAKAN KOPERASI"), (hk, "KET")):
        ws[f"{h}{header}"] = t
    for c in ws[header]:
        c.font = TEBAL
        c.fill = ABU
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[header].height = 42
    kode_kolom = {}
    for h, k in kb["peta"].items():
        kode_kolom.setdefault(k, h)                    # satu kode, satu kolom (kolom pertama)

    def tulis(r, x):
        ws.cell(r, 1, x.get("no") or None)
        ws.cell(r, 2, x["nama"])
        ws.cell(r, 3, x.get("pangkat") or "")
        ws.cell(r, 5, x.get("nrp_asli") or "")
        if x.get("gaji") is not None:
            ws.cell(r, 6, int(x["gaji"])).number_format = RP
        if x["total"] is not None:
            for k, v in x["potongan"].items():
                h = kode_kolom.get(k)
                if h:
                    ws[f"{h}{r}"] = v
                    ws[f"{h}{r}"].number_format = RP
            ws[f"{hj}{r}"] = x["total"]
            ws[f"{hy}{r}"] = x["ybs"]
            for h in (hj, hy):
                ws[f"{h}{r}"].number_format = RP
            if x["ybs"] < 0:
                ws[f"{hy}{r}"].font = MERAH
            if x["kop"]["tunggakan"]:
                ws[f"{ht}{r}"] = x["kop"]["tunggakan"]
                ws[f"{ht}{r}"].number_format = RP
                ws[f"{ht}{r}"].fill = KUNING
        if x["ket"]:
            ws[f"{hk}{r}"] = ", ".join(dict.fromkeys(x["ket"]))
            ws[f"{hk}{r}"].font = MERAH if {"YBS MINUS", "TIDAK ADA DI DATABASE"} & set(x["ket"]) else Font()

    for x in data:
        tulis(x["baris"], x)
    hilang = [k for k in {kk for x in kel["baris"] for kk in x["potongan"]} if k not in kode_kolom]
    rj = max([baris_jumlah] + [x["baris"] + 1 for x in data])
    ws.cell(rj, 2, "JUMLAH").font = TEBAL
    hitung = [x for x in data if x["total"] is not None]
    ws.cell(rj, 6, sum(int(x["gaji"]) for x in hitung)).number_format = RP
    for h in kolom:
        k = kb["peta"].get(h)
        if k and kode_kolom.get(k) == h:
            ws[f"{h}{rj}"] = sum(x["potongan"].get(k, 0) for x in hitung)
            ws[f"{h}{rj}"].number_format = RP
    ws[f"{hj}{rj}"] = sum(x["total"] for x in hitung)
    ws[f"{hy}{rj}"] = sum(x["ybs"] for x in hitung)
    ws[f"{ht}{rj}"] = sum(x["kop"]["tunggakan"] for x in hitung)
    for h in (hj, hy, ht):
        ws[f"{h}{rj}"].number_format = RP
    for c in ws[rj]:
        c.font = TEBAL
    r = rj + 2
    if hilang:
        ws.cell(r, 2, "PERHATIAN: buku ini tidak punya kolom untuk " + ", ".join(LABEL[k] for k in hilang)
                + "; potongan itu ikut dihitung di JUMLAH POTONGAN tetapi tidak punya kolom sendiri.").font = MERAH
        r += 2
    if lain:
        ws.cell(r, 2, "ADA DI DPP, BELUM ADA DI BUKU (tambahkan barisnya di buku dulu)").font = TEBAL
        r += 1
        for x in lain:
            tulis(r, x)
            r += 1
    ws.freeze_panes = f"G{header + 1}"
    _lebar(ws, {"A": 6, "B": 30, "C": 16, "D": 4, "E": 20, "F": 13, hk: 28})
    for h in kolom + [hj, hy, ht]:
        ws.column_dimensions[h].width = 12


def _sheet_tunggakan(wb, hasil, kop_atas):
    ws = wb.create_sheet("TUNGGAKAN")
    r = 1
    for t in kop_atas:
        ws.cell(r, 1, t)
        r += 1
    r += 1
    ws.cell(r, 1, "DAFTAR TUNGGAKAN ANGGOTA").font = TEBAL
    ws.cell(r + 1, 1, f"BULAN : {hasil['bulan']}")
    r += 3
    for j, t in enumerate(["NO", "N A M A", "PANGKAT/", "GJ YG DITERIMA", "TAGIHAN KOPERASI", "POT SIMP WAJIB",
                           "POT PINJAMAN", "TERPOTONG", "TUNGGAKAN", "KET"], 1):
        ws.cell(r, j, t).font = TEBAL
    for j, t in enumerate(["URT", "", "GOL", "( Rp )", "( Rp )", "( Rp )", "( Rp )", "( Rp )", "( Rp )", ""], 1):
        ws.cell(r + 1, j, t).font = TEBAL
    for j in range(1, 11):
        ws.cell(r + 2, j, j)
    r += 3
    n, jml = 0, [0] * 5
    for kel in hasil["kelompok"]:
        for x in kel["baris"]:
            if not x["kop"] or not x["kop"]["tunggakan"]:
                continue
            n += 1
            k = x["kop"]
            nilai = [k["tagihan"], k["sw"], k["kp"], k["terpotong"], k["tunggakan"]]
            jml = [a + b for a, b in zip(jml, nilai)]
            ws.cell(r, 1, n)
            ws.cell(r, 2, x["nama"])
            ws.cell(r, 3, x.get("pangkat") or "")
            ws.cell(r, 4, int(x["gaji"])).number_format = RP
            for j, v in enumerate(nilai, 5):
                ws.cell(r, j, v).number_format = RP
            ws.cell(r, 10, kel["jenis"])
            r += 1
    ws.cell(r, 2, "JUMLAH").font = TEBAL
    for j, v in enumerate(jml, 5):
        c = ws.cell(r, j, v)
        c.number_format, c.font = RP, TEBAL
    if not n:
        ws.cell(r + 2, 2, "Tidak ada tunggakan bulan ini.")
    _lebar(ws, {"A": 6, "B": 30, "C": 16, "D": 15, "E": 16, "F": 15, "G": 15, "H": 15, "I": 15, "J": 8})
    return n


def _sheet_mass_debet(wb, hasil, bank):
    ws = wb.create_sheet("MASS DEBET")
    baris = [x for kel in hasil["kelompok"] for x in kel["baris"] if x["total"] is not None and x["mass_debet"] > 0]
    ws.append(["Rekening Debet", None, None, "Jumlah amount", "Jumlah charge", "jumlah record"])
    ws.append([None, "TTP POTONGAN SATUAN DITAJENAD", "POT GAJI", sum(x["mass_debet"] for x in baris), 0, len(baris)])
    ws.append(["Rekening Debet", "Nama", "Remarks", "Amount", "Charge"])
    tanpa = 0
    for i, x in enumerate(baris, 1):
        rek = x.get("rekening") or ""
        tanpa += not rek
        ws.append([rek, x.get("nama_rekening") or x["nama"], None, x["mass_debet"], i])
        ws.cell(ws.max_row, 1).number_format = "@"
    for r in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in r:
            c.number_format = RP
    _lebar(ws, {"A": 20, "B": 34, "C": 10, "D": 15, "E": 10, "F": 13})
    return {"baris": len(baris), "total": sum(x["mass_debet"] for x in baris), "tanpa_rekening": tanpa}


def tulis_daftar(path, hasil, kolom: dict, kop_atas=()) -> dict:
    """Tulis berkas Daftar Potongan: POTONGAN TNI/PNS (baris dan kolom sama dengan buku), TUNGGAKAN, MASS
    DEBET. kolom = {sheet: {"kb": kolom_buku(), "header": n, "baris_jumlah": n}}."""
    wb = Workbook()
    wb.remove(wb.active)
    for kel in hasil["kelompok"]:
        k = kolom[kel["sheet"]]
        _sheet_potongan(wb, kel, k["kb"], k["header"], k["baris_jumlah"], hasil["bank"], hasil["bulan"])
    n_tung = _sheet_tunggakan(wb, hasil, kop_atas)
    md = _sheet_mass_debet(wb, hasil, hasil["bank"])
    wb.save(path)
    return {"tunggakan": n_tung, "mass_debet": md}


def _kop_tunggakan(sumber) -> list:
    """Dua baris kop (MARKAS BESAR ... / DIREKTORAT ...) dari sheet TUNGGAKAN buku sumber, bila ada."""
    try:
        wb = load_workbook(sumber, read_only=True, data_only=True)
    except Exception:                                   # noqa: BLE001
        return []
    try:
        if "TUNGGAKAN" not in wb.sheetnames:
            return []
        out = []
        for r in wb["TUNGGAKAN"].iter_rows(min_row=1, max_row=5, max_col=1, values_only=True):
            t = str(r[0] or "").strip()
            if t and "TUNGGAKAN" not in t.upper() and "BULAN" not in t.upper():
                out.append(t)
        return out[:3]
    finally:
        wb.close()


def daftar_proses(snap, sumber, tpl, buku, update, dpp_tanpa_buku, bulan_baru, bank, tujuan, log=print) -> dict:
    """Dipanggil jalankan.proses(): susun baris TNI (buku + gaji DPP), PNS (BUKU PNS + gaji buku lalu,
    sampai pembaca DPP PNS ada), dan anggota DPP yang belum ada di buku; hitung; tulis berkas."""
    rows_tni = _rows(sumber, tpl["sheet"], buku["baris_jumlah"])
    kb_tni = kolom_buku(rows_tni, buku["header"])
    tni = [{"baris": b["baris"], "no": b["no"], "nama": b["nama"], "pangkat": b["pangkat"], "nrp": b["nrp"],
            "nrp_asli": b["nrp_asli"], "rekening": b["rekening"], "nama_rekening": b.get("nama_rekening", ""),
            "gaji": update.get(b["baris"])} for b in buku["anggota"]]
    tni += [{"baris": None, "no": None, "nama": d["nama"], "pangkat": d.get("pangkat", ""), "nrp": digit(d["nrp"]),
             "nrp_asli": str(d["nrp"]), "rekening": digit(d.get("rekening", "")), "nama_rekening": "",
             "gaji": d["dibayarkan"]} for d in dpp_tanpa_buku]
    kelompok = [{"sheet": tpl["sheet"], "jenis": "TNI", "baris": tni}]
    kolom = {tpl["sheet"]: {"kb": kb_tni, "header": buku["header"], "baris_jumlah": buku["baris_jumlah"]}}
    catatan_pns = ""
    try:
        wb = load_workbook(sumber, read_only=True)
        ada_pns = "BUKU PNS" in wb.sheetnames
        wb.close()
        sb = _sheet_buku(sumber, tpl, "BUKU PNS", "PNS") if ada_pns else None
    except SystemExit as e:
        sb, catatan_pns = None, str(e.code)
    if sb and sb[1]["anggota"]:
        _, bp, rows_pns = sb
        pns = [{"baris": b["baris"], "no": b["no"], "nama": b["nama"], "pangkat": b["pangkat"], "nrp": b["nrp"],
                "nrp_asli": b["nrp_asli"], "rekening": b["rekening"], "nama_rekening": b.get("nama_rekening", ""),
                "gaji": (int(round(b["gaji_lama"])) if isinstance(b["gaji_lama"], (int, float)) and b["gaji_lama"]
                         else None)} for b in bp["anggota"]]
        kelompok.append({"sheet": "BUKU PNS", "jenis": "PNS", "baris": pns})
        kolom["BUKU PNS"] = {"kb": kolom_buku(rows_pns, bp["header"]), "header": bp["header"],
                             "baris_jumlah": bp["baris_jumlah"]}
    hasil = hitung_daftar(snap, bulan_baru, bank, kelompok)
    if len(kelompok) > 1:
        hasil["peringatan"].insert(0, {"tingkat": "INFO", "kategori": "Gaji PNS dari BUKU PNS bulan lalu",
                                       "detail": "Pembaca DPP PNS belum ada, jadi gaji PNS diambil dari kolom GAJI "
                                                 "BUKU PNS buku sumber. Bila gaji PNS bulan ini berubah, hitungan "
                                                 "kurang gaji PNS bisa meleset.", "nama": "", "nrp": "", "letak": ""})
    elif catatan_pns:
        hasil["peringatan"].insert(0, {"tingkat": "SEDANG", "kategori": "BUKU PNS tidak terbaca",
                                       "detail": catatan_pns, "nama": "", "nrp": "", "letak": ""})
    meta = (snap["aturan"].get("arisan") or {}).get(bank or "")
    if not meta:
        hasil["peringatan"].insert(0, {"tingkat": "SEDANG", "kategori": "Daftar arisan belum diimpor",
                                       "detail": f"Belum pernah ada DAFTAR ARISAN PERSIT {bank or ''} yang diimpor, "
                                                 "jadi potongan arisan memakai angka yang ada di database (bila ada). "
                                                 "Impor lewat Data Personel > Impor.", "nama": "", "nrp": "",
                                       "letak": ""})
    else:
        hasil["arisan"] = meta
    info = tulis_daftar(tujuan, hasil, kolom, _kop_tunggakan(sumber))
    hasil.update(info)
    t = [x for k in hasil["kelompok"] for x in k["baris"] if x["total"] is not None]
    hasil["ringkas"] = {"dihitung": len(t), "total": sum(x["total"] for x in t),
                        "tunggakan": sum(x["kop"]["tunggakan"] for x in t),
                        "orang_tunggakan": info["tunggakan"],
                        "minus": sum(1 for x in t if x["ybs"] < 0),
                        "mass_debet": info["mass_debet"]["total"], "mass_debet_baris": info["mass_debet"]["baris"],
                        "tanpa_rekening": info["mass_debet"]["tanpa_rekening"],
                        "per_sheet": {k["sheet"]: sum(1 for x in k["baris"] if x["total"] is not None)
                                      for k in hasil["kelompok"]}}
    md = info["mass_debet"]
    log(f"      Daftar potongan: {len(t)} orang dihitung, potongan Rp {_rp(hasil['ringkas']['total']) or '0'}, "
        f"MASS DEBET {md['baris']} baris Rp {_rp(md['total']) or '0'}, tunggakan {info['tunggakan']} orang"
        + (f", {hasil['ringkas']['minus']} YBS minus" if hasil["ringkas"]["minus"] else ""))
    if md["tanpa_rekening"]:
        hasil["peringatan"].append({"tingkat": "SEDANG", "kategori": "MASS DEBET tanpa nomor rekening",
                                    "detail": f"{md['tanpa_rekening']} baris MASS DEBET tidak punya nomor rekening. "
                                              "Lengkapi rekeningnya di Data Personel.", "nama": "", "nrp": "",
                                    "letak": ""})
    return hasil
