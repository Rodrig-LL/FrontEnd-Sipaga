"""
Database personel (master personel): acuan pencocokan DPP, disimpan permanen di data/personel.db.

Dulu acuan pencocokan adalah berkas nominatif yang diseret setiap proses. Juru bayar mengoreksinya
16 Okt 2026: nominatif "native ada di app ... udah ada di database nya dan bisa crud nya".
Keputusan rinciannya ada di bagian 4h dokumen konteks; ringkasnya:
- satu berkas SQLite (pustaka bawaan) di folder aplikasi, supaya ikut terbawa bila folder disalin;
- isi awal dari Excel nominatif (nominatif.baca_nominatif), bank dan rekening dari pembukuan;
- impor ulang = gabung dengan pratinjau: NRP baru ditambah, yang ada diperbarui (bank dan rekening
  tidak disentuh), yang tidak tercantum lagi hanya DITANDAI untuk dicek;
- yang pindah/pensiun ditandai keluar, bukan dihapus; hapus hanya untuk data salah input;
- setiap perubahan dicatat di tabel riwayat yang tidak bisa diubah atau dihapus (dijaga trigger).
  Tidak ada login, jadi "siapa" = nama operator dari halaman Pengaturan.
"""
import json
import re
import sqlite3
import threading
from contextlib import contextmanager
from datetime import date, datetime
from pathlib import Path

from nominatif import baca_nominatif
from pembukuan import AMBANG_NAMA, _kotak, baca_buku, digit, mirip, muat_template, pastikan_xlsx

DASAR = Path(__file__).resolve().parent
DB_BAWAAN = DASAR / "data" / "personel.db"
JENIS = ("TNI", "PNS")
BANK = ("BRI", "BSI", "BTN")
MIN_DIGIT = 5
# Kolom yang diisi dari nominatif. Bank, rekening, status, dan tanda sengaja tidak: impor nominatif
# tidak boleh menimpa yang sudah dilengkapi juru bayar.
DARI_NOMINATIF = ("jenis", "nama", "pangkat", "nrp_asli", "jabatan", "ket", "seksi", "tgl_lahir")
DARI_BUKU = ("bank", "rekening", "nama_rekening")
KOLOM = ("jenis", "nama", "pangkat", "nrp", "nrp_asli", "jabatan", "ket", "seksi", "tgl_lahir",
         "bank", "rekening", "nama_rekening", "status", "tgl_keluar", "alasan_keluar", "tandai")
LABEL = {"jenis": "jenis", "nama": "nama", "pangkat": "pangkat/golongan", "nrp": "NRP/NIP",
         "nrp_asli": "NRP/NIP", "jabatan": "jabatan", "ket": "KET", "seksi": "seksi",
         "tgl_lahir": "tanggal lahir", "bank": "bank", "rekening": "no rekening",
         "nama_rekening": "nama rekening", "status": "status", "tgl_keluar": "tanggal keluar",
         "alasan_keluar": "alasan keluar", "tandai": "tanda perlu dicek"}
RE_BANK_BERKAS = re.compile(r"(?<![A-Z])(BRI|BSI|BTN)(?![A-Z])")

SKEMA = """
CREATE TABLE IF NOT EXISTS personel (
  id INTEGER PRIMARY KEY,
  jenis TEXT NOT NULL CHECK (jenis IN ('TNI', 'PNS')),
  nama TEXT NOT NULL,
  pangkat TEXT NOT NULL DEFAULT '',
  nrp TEXT NOT NULL UNIQUE,
  nrp_asli TEXT NOT NULL DEFAULT '',
  jabatan TEXT NOT NULL DEFAULT '',
  ket TEXT NOT NULL DEFAULT '',
  seksi TEXT NOT NULL DEFAULT '',
  tgl_lahir TEXT NOT NULL DEFAULT '',
  bank TEXT NOT NULL DEFAULT '' CHECK (bank IN ('', 'BRI', 'BSI', 'BTN')),
  rekening TEXT NOT NULL DEFAULT '',
  nama_rekening TEXT NOT NULL DEFAULT '',
  status TEXT NOT NULL DEFAULT 'aktif' CHECK (status IN ('aktif', 'keluar')),
  tgl_keluar TEXT NOT NULL DEFAULT '',
  alasan_keluar TEXT NOT NULL DEFAULT '',
  tandai TEXT NOT NULL DEFAULT '',
  dibuat TEXT NOT NULL,
  diubah TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS riwayat (
  id INTEGER PRIMARY KEY,
  waktu TEXT NOT NULL,
  operator TEXT NOT NULL,
  aksi TEXT NOT NULL,
  personel_id INTEGER,
  nrp TEXT NOT NULL DEFAULT '',
  nama TEXT NOT NULL DEFAULT '',
  perubahan TEXT NOT NULL DEFAULT '{}'
);
CREATE INDEX IF NOT EXISTS riwayat_personel ON riwayat (personel_id);
CREATE TABLE IF NOT EXISTS meta (kunci TEXT PRIMARY KEY, nilai TEXT NOT NULL);
INSERT OR IGNORE INTO meta VALUES ('versi', '0');
CREATE TRIGGER IF NOT EXISTS riwayat_tidak_diubah BEFORE UPDATE ON riwayat
  BEGIN SELECT RAISE(ABORT, 'riwayat tidak boleh diubah'); END;
CREATE TRIGGER IF NOT EXISTS riwayat_tidak_dihapus BEFORE DELETE ON riwayat
  BEGIN SELECT RAISE(ABORT, 'riwayat tidak boleh dihapus'); END;
-- Potongan satuan (potongan.py, bagian 4i dokumen konteks). Kode jenis ada di potongan.JENIS.
CREATE TABLE IF NOT EXISTS potongan (
  personel_id INTEGER NOT NULL,
  kode TEXT NOT NULL,
  nominal INTEGER NOT NULL CHECK (nominal > 0),
  PRIMARY KEY (personel_id, kode)
);
CREATE TABLE IF NOT EXISTS pinjaman (
  id INTEGER PRIMARY KEY,
  personel_id INTEGER NOT NULL,
  pemberi TEXT NOT NULL CHECK (pemberi IN ('KOPERASI', 'LURJA')),
  cicilan INTEGER NOT NULL CHECK (cicilan > 0),
  mulai TEXT NOT NULL,
  tenor INTEGER CHECK (tenor IS NULL OR tenor > 0),
  tgl_cair TEXT NOT NULL DEFAULT '',
  catatan TEXT NOT NULL DEFAULT '',
  status TEXT NOT NULL DEFAULT 'aktif' CHECK (status IN ('aktif', 'lunas', 'batal')),
  dibuat TEXT NOT NULL,
  diubah TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS pinjaman_personel ON pinjaman (personel_id);
CREATE TABLE IF NOT EXISTS pokok (
  personel_id INTEGER PRIMARY KEY,
  bulan TEXT NOT NULL,
  nominal INTEGER NOT NULL CHECK (nominal > 0)
);
CREATE TABLE IF NOT EXISTS aturan (kunci TEXT PRIMARY KEY, nilai TEXT NOT NULL);
"""

_path = DB_BAWAAN
_tulis = threading.Lock()             # server berutas banyak: satu penulis pada satu waktu


def atur_path(p):
    """Pindahkan database (dipakai uji dan flag --db aplikasi.py)."""
    global _path
    _path = Path(p)


def lokasi() -> Path:
    return _path


@contextmanager
def _sambung():
    """Satu koneksi per pemakaian: commit bila lancar, rollback bila galat, lalu SELALU ditutup.
    (Konteks sqlite3 bawaan hanya commit, tidak menutup; di Windows berkasnya bisa tetap terkunci.)"""
    _path.parent.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(_path, timeout=15)
    try:
        con.row_factory = sqlite3.Row
        con.executescript(SKEMA)
        with con:
            yield con
    finally:
        con.close()


def _sekarang():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _versi(con) -> int:
    return int(con.execute("SELECT nilai FROM meta WHERE kunci = 'versi'").fetchone()[0])


def _naik(con):
    con.execute("UPDATE meta SET nilai = CAST(nilai AS INTEGER) + 1 WHERE kunci = 'versi'")


def _catat(con, aksi, operator, pid, nrp, nama, perubahan):
    con.execute("INSERT INTO riwayat (waktu, operator, aksi, personel_id, nrp, nama, perubahan) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (_sekarang(), (operator or "").strip() or "(operator belum diisi)", aksi, pid, nrp or "",
                 nama or "", json.dumps(perubahan, ensure_ascii=False)))


def _tgl(v) -> str:
    """Tanggal dari Excel atau teks 'YYYY-MM-DD' -> 'YYYY-MM-DD'; selain itu kosong."""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()[:10]
    try:
        return datetime.strptime(s, "%Y-%m-%d").strftime("%Y-%m-%d") if s else ""
    except ValueError:
        return ""


# ------------------------------------------------------------------------------- baca
def daftar() -> list:
    with _sambung() as con:
        rows = con.execute("SELECT * FROM personel ORDER BY jenis DESC, nama").fetchall()
    return [dict(r) for r in rows]


def ringkas() -> dict:
    with _sambung() as con:
        q = lambda sql: con.execute(sql).fetchone()[0]   # noqa: E731
        return {"total": q("SELECT COUNT(*) FROM personel"),
                "tni_aktif": q("SELECT COUNT(*) FROM personel WHERE jenis='TNI' AND status='aktif'"),
                "pns_aktif": q("SELECT COUNT(*) FROM personel WHERE jenis='PNS' AND status='aktif'"),
                "keluar": q("SELECT COUNT(*) FROM personel WHERE status='keluar'"),
                "tandai": q("SELECT COUNT(*) FROM personel WHERE tandai<>''"),
                "tanpa_bank": q("SELECT COUNT(*) FROM personel WHERE status='aktif' AND bank=''"),
                "diubah": q("SELECT COALESCE(MAX(diubah), '') FROM personel"),
                "versi": _versi(con)}


def ambil(pid):
    with _sambung() as con:
        r = con.execute("SELECT * FROM personel WHERE id = ?", (pid,)).fetchone()
    return dict(r) if r else None


def riwayat(pid=None, batas=2000) -> list:
    with _sambung() as con:
        if pid is None:
            rows = con.execute("SELECT * FROM riwayat ORDER BY id DESC LIMIT ?", (batas,)).fetchall()
        else:
            rows = con.execute("SELECT * FROM riwayat WHERE personel_id = ? ORDER BY id DESC LIMIT ?",
                               (pid, batas)).fetchall()
    return [{**dict(r), "perubahan": json.loads(r["perubahan"])} for r in rows]


def acuan(jenis="TNI") -> dict:
    """Personel satu jenis untuk pencocokan DPP (nominatif.periksa), aktif maupun keluar."""
    orang = [p for p in daftar() if p["jenis"] == jenis]
    r = ringkas()
    return {"orang": orang, "jenis": jenis, "diubah": r["diubah"],
            "lain": r["pns_aktif"] if jenis == "TNI" else r["tni_aktif"]}


# ------------------------------------------------------------------------------- tulis
def _bersihkan(data: dict) -> dict:
    v = {k: str(data.get(k) if data.get(k) is not None else "").strip() for k in KOLOM}
    v["jenis"] = v["jenis"].upper()
    v["bank"] = v["bank"].upper()
    v["status"] = v["status"].lower() or "aktif"
    v["nrp"] = digit(v["nrp_asli"] or v["nrp"])
    v["nrp_asli"] = v["nrp_asli"] or v["nrp"]
    v["rekening"] = re.sub(r"[\s.\-]", "", v["rekening"])
    if v["status"] == "aktif":
        v["tgl_keluar"] = v["alasan_keluar"] = ""
    return v


def _periksa(con, v, pid=None) -> dict:
    g = {}
    if v["jenis"] not in JENIS:
        g["jenis"] = "Pilih TNI atau PNS."
    if not v["nama"]:
        g["nama"] = "Nama harus diisi."
    if len(v["nrp"]) < MIN_DIGIT:
        g["nrp_asli"] = f"{'NIP' if v['jenis'] == 'PNS' else 'NRP'} harus berisi minimal {MIN_DIGIT} angka."
    else:
        r = con.execute("SELECT id, nama FROM personel WHERE nrp = ? AND id IS NOT ?", (v["nrp"], pid)).fetchone()
        if r:
            g["nrp_asli"] = f"Nomor ini sudah dipakai oleh {r['nama']}."
    if v["bank"] not in ("",) + BANK:
        g["bank"] = "Pilih BRI, BSI, BTN, atau kosongkan."
    if v["rekening"] and not v["rekening"].isdigit():
        g["rekening"] = "Nomor rekening hanya boleh berisi angka."
    if v["tgl_lahir"] and not _tgl(v["tgl_lahir"]):
        g["tgl_lahir"] = "Tanggal lahir tidak sah."
    if v["status"] not in ("aktif", "keluar"):
        g["status"] = "Status harus aktif atau keluar."
    elif v["status"] == "keluar":
        if not _tgl(v["tgl_keluar"]):
            g["tgl_keluar"] = "Tanggal keluar harus diisi."
        if not v["alasan_keluar"]:
            g["alasan_keluar"] = "Alasan keluar harus diisi (mis. pindah satuan, pensiun)."
    return g


def simpan(data: dict, operator: str, pid=None) -> dict:
    """Tambah (pid None) atau ubah satu personel. Mengembalikan {"ok", "id"} atau {"ok": False,
    "galat": {kolom: pesan}}. Perubahan dicatat di riwayat, kolom demi kolom."""
    v = _bersihkan(data)
    with _tulis, _sambung() as con:
        g = _periksa(con, v, pid)
        if g:
            return {"ok": False, "galat": g}
        v["tgl_lahir"], v["tgl_keluar"] = _tgl(v["tgl_lahir"]), _tgl(v["tgl_keluar"])
        now = _sekarang()
        if pid is None:
            cur = con.execute(f"INSERT INTO personel ({', '.join(KOLOM)}, dibuat, diubah) VALUES "
                              f"({', '.join('?' * len(KOLOM))}, ?, ?)", [v[k] for k in KOLOM] + [now, now])
            pid = cur.lastrowid
            _catat(con, "tambah", operator, pid, v["nrp"], v["nama"],
                   {k: ["", v[k]] for k in KOLOM if v[k] and k != "nrp"})
        else:
            lama = con.execute("SELECT * FROM personel WHERE id = ?", (pid,)).fetchone()
            if not lama:
                return {"ok": False, "galat": {"_": "Personel ini sudah tidak ada (mungkin dihapus di tab lain)."}}
            beda = {k: [lama[k], v[k]] for k in KOLOM if lama[k] != v[k] and k != "nrp"}
            if not beda:
                return {"ok": True, "id": pid, "berubah": False}
            con.execute(f"UPDATE personel SET {', '.join(k + ' = ?' for k in KOLOM)}, diubah = ? WHERE id = ?",
                        [v[k] for k in KOLOM] + [now, pid])
            aksi = ("keluar" if lama["status"] == "aktif" and v["status"] == "keluar" else
                    "aktif kembali" if lama["status"] == "keluar" and v["status"] == "aktif" else "ubah")
            _catat(con, aksi, operator, pid, v["nrp"], v["nama"], beda)
        _naik(con)
    return {"ok": True, "id": pid, "berubah": True}


def hapus(pid, operator) -> dict:
    """Hapus permanen (untuk data salah input). Isi terakhirnya disimpan di riwayat."""
    with _tulis, _sambung() as con:
        lama = con.execute("SELECT * FROM personel WHERE id = ?", (pid,)).fetchone()
        if not lama:
            return {"ok": False, "galat": {"_": "Personel ini sudah tidak ada."}}
        # Potongan, pinjaman, dan simpanan pokoknya ikut terhapus; ringkasannya ikut di riwayat.
        ikut = {}
        pot = con.execute("SELECT kode, nominal FROM potongan WHERE personel_id = ?", (pid,)).fetchall()
        if pot:
            ikut["potongan"] = [", ".join(f"{r['kode']} {r['nominal']}" for r in pot), ""]
        pj = con.execute("SELECT pemberi, cicilan, status FROM pinjaman WHERE personel_id = ?", (pid,)).fetchall()
        if pj:
            ikut["pinjaman"] = [", ".join(f"{r['pemberi']} {r['cicilan']}/bln ({r['status']})" for r in pj), ""]
        for tabel in ("potongan", "pinjaman", "pokok"):
            con.execute(f"DELETE FROM {tabel} WHERE personel_id = ?", (pid,))
        con.execute("DELETE FROM personel WHERE id = ?", (pid,))
        _catat(con, "hapus", operator, pid, lama["nrp"], lama["nama"],
               {**{k: [lama[k], ""] for k in KOLOM if lama[k] and k != "nrp"}, **ikut})
        _naik(con)
    return {"ok": True}


# ------------------------------------------------------------------------------- impor
def rencana_nominatif(paths) -> dict:
    """Pratinjau impor nominatif (militer dan/atau PNS, satu berkas per jenis). Belum menulis apa pun."""
    berkas = [baca_nominatif(p) for p in paths]
    per_jenis = {}
    for b in berkas:
        per_jenis.setdefault("PNS" if b["jenis"] == "pns" else "TNI", []).append(b)
    ganda = {j: bs for j, bs in per_jenis.items() if len(bs) > 1}
    if ganda:
        raise SystemExit(_kotak(
            "lebih dari satu nominatif untuk jenis yang sama",
            [f"  - {j}: " + ", ".join(b["file"] for b in bs) for j, bs in ganda.items()],
            "Program tidak bisa memilih mana yang berlaku. Sisakan satu per jenis, yang terbaru."))
    with _sambung() as con:
        versi = _versi(con)
        ada = {r["nrp"]: dict(r) for r in con.execute("SELECT * FROM personel")}
    hasil = {"jenis_impor": "nominatif", "versi": versi, "tambah": [], "ubah": [], "tetap": 0,
             "hilang": [], "masalah": [], "catatan": [],
             "berkas": [{"file": b["file"], "sheet": b["sheet"], "judul": b["judul"],
                         "jenis": "PNS" if b["jenis"] == "pns" else "TNI", "jumlah": len(b["orang"])}
                        for b in berkas]}
    tanggal = datetime.now().strftime("%d-%m-%Y")
    for jenis, (b,) in per_jenis.items():
        dilihat = {}
        for o in b["orang"]:
            baris = {"jenis": jenis, "nama": o["nama"], "pangkat": o["pangkat"], "nrp": o["nrp"],
                     "nrp_asli": o["nrp_asli"], "jabatan": o["jabatan"], "ket": o["ket"],
                     "seksi": o["seksi"], "tgl_lahir": _tgl(o["tgl_lahir"])}
            if o["nrp"] in dilihat:
                hasil["masalah"].append({"file": b["file"], "baris": o["baris"], "nama": o["nama"],
                                         "nrp": o["nrp_asli"],
                                         "alasan": f"NRP/NIP ganda di berkas (juga baris {dilihat[o['nrp']]}); "
                                                   "hanya yang pertama dipakai."})
                continue
            if not o["nama"]:
                hasil["masalah"].append({"file": b["file"], "baris": o["baris"], "nama": "", "nrp": o["nrp_asli"],
                                         "alasan": "Nama kosong; baris dilewati."})
                continue
            dilihat[o["nrp"]] = o["baris"]
            lama = ada.get(o["nrp"])
            if not lama:
                hasil["tambah"].append(baris)
                continue
            beda = {k: [lama[k], baris[k]] for k in DARI_NOMINATIF if lama[k] != baris[k]}
            if lama["tandai"]:
                beda["tandai"] = [lama["tandai"], ""]
            if lama["status"] == "keluar":
                hasil["catatan"].append(f"{o['nama']} ({o['nrp_asli']}) tercatat KELUAR di database, tetapi "
                                        f"tercantum di {b['file']}. Statusnya tidak diubah otomatis; periksa.")
            if beda:
                hasil["ubah"].append({"id": lama["id"], "nama": lama["nama"], "nrp": lama["nrp_asli"],
                                      "perubahan": beda})
            else:
                hasil["tetap"] += 1
        for p in ada.values():
            if p["jenis"] == jenis and p["status"] == "aktif" and p["nrp"] not in dilihat:
                tanda = f"Tidak tercantum di {b['judul'] or b['file']} (impor {tanggal})."
                hasil["hilang"].append({"id": p["id"], "nama": p["nama"], "nrp": p["nrp_asli"],
                                        "pangkat": p["pangkat"], "bank": p["bank"], "tandai": tanda})
    return hasil


def _tpl_sheet(tpl, sheet, ws_rows, header_cari=(1, 19)):
    """Peta kolom untuk satu sheet buku: rekening = kolom berjudul NO REKENING bila ada (BUKU PNS:
    AQ di BRI/BTN, AP di BSI), selain itu peta template_pembukuan.json (BUKU TNI: AS). Nama
    rekening selalu kolom sesudahnya - terukur pada keenam sheet buku asli, 16 Okt 2026."""
    from openpyxl.utils import get_column_letter
    t = json.loads(json.dumps({k: v for k, v in tpl.items() if not k.startswith("_")}))
    t["_sumber"], t["_catatan"] = tpl.get("_sumber", ""), []
    t["sheet"] = sheet
    for r in ws_rows[:header_cari[1] + 3]:
        for j, v in enumerate(r):
            if re.sub(r"\s+", "", str(v or "")).upper() == "NOREKENING":
                t["kolom"]["rekening"] = get_column_letter(j + 1)
                t["kolom"]["nama_rekening"] = get_column_letter(j + 2)
                return t
    if sheet != tpl["sheet"]:
        return None                        # sheet lain tanpa judul NO REKENING: tidak ditebak
    return t


def rencana_bank(paths) -> dict:
    """Pratinjau pengisian bank dan rekening dari pembukuan (BUKU TNI dan BUKU PNS tiap bank).
    Bank diambil dari nama berkas (... BRI ..., ... BSI ..., ... BTN ...)."""
    import openpyxl
    tpl = muat_template()
    with _sambung() as con:
        versi = _versi(con)
        semua = [dict(r) for r in con.execute("SELECT * FROM personel")]
    hasil = {"jenis_impor": "bank", "versi": versi, "ubah": [], "tetap": 0, "tidak_ketemu": [],
             "masalah": [], "berkas": [], "catatan": []}
    usul = {}                              # id personel -> (bank, rekening, nama_rekening, asal)
    for path in paths:
        path = Path(path)
        m = RE_BANK_BERKAS.search(path.stem.upper())
        if not m:
            raise SystemExit(_kotak(
                "bank pembukuan tidak diketahui",
                [f"Berkas : {path.name}",
                 "Nama berkasnya tidak menyebut BRI, BSI, atau BTN, padahal dari situlah program tahu",
                 "bank para personel di dalamnya."],
                "Ganti nama berkasnya (mis. PEMBUKUAN GAJI BRI BLN OKTOBER 2026.xlsx), lalu coba lagi."))
        bank = m.group(1)
        sumber = pastikan_xlsx(path)
        wb = openpyxl.load_workbook(sumber, read_only=True, data_only=True)
        try:
            sheet_rows = {s: list(wb[s].iter_rows(min_row=1, max_row=25, values_only=True))
                          for s in ("BUKU TNI", "BUKU PNS") if s in wb.sheetnames}
        finally:
            wb.close()
        info = {"file": path.name, "bank": bank, "sheet": []}
        for sheet, jenis in (("BUKU TNI", "TNI"), ("BUKU PNS", "PNS")):
            if sheet not in sheet_rows:
                continue
            t = _tpl_sheet(tpl, sheet, sheet_rows[sheet])
            if t is None:
                hasil["catatan"].append(f"{path.name}, sheet {sheet}: kolom NO REKENING tidak ditemukan; dilewati.")
                continue
            if jenis == "PNS":
                t["akhir_tabel"] = {**t["akhir_tabel"], "kolom": "B"}     # BUKU PNS: JUMLAH di kolom B
            buku = baca_buku(sumber, t, boleh_kosong=True)["anggota"]
            info["sheet"].append({"sheet": sheet, "baris": len(buku), "rekening": t["kolom"]["rekening"]})
            calon = [p for p in semua if p["jenis"] == jenis]
            per_nrp = {p["nrp"]: p for p in calon}
            for b in buku:
                p, cara = per_nrp.get(b["nrp"]), "NRP"
                if not p:
                    skor = sorted(((mirip(b["nama"], q["nama"]), q["id"], q) for q in calon), key=lambda x: (-x[0], x[1]))
                    if skor and skor[0][0] >= AMBANG_NAMA and (len(skor) == 1 or skor[1][0] < skor[0][0]):
                        p, cara = skor[0][2], "nama"
                if not p:
                    hasil["tidak_ketemu"].append({"file": path.name, "sheet": sheet, "baris": b["baris"],
                                                  "nama": b["nama"], "nrp": b["nrp_asli"], "bank": bank})
                    continue
                usul.setdefault(p["id"], []).append({"bank": bank, "rekening": b["rekening"],
                                                     "nama_rekening": b["nama_rekening"], "cara": cara,
                                                     "asal": f"{path.name} {sheet} baris {b['baris']}"})
        hasil["berkas"].append(info)
    per_id = {p["id"]: p for p in semua}
    for pid, us in usul.items():
        p = per_id[pid]
        if len(us) > 1:                    # lebih dari satu baris buku menunjuk orang yang sama
            hasil["masalah"].append({"id": pid, "nama": p["nama"], "nrp": p["nrp_asli"],
                                     "alasan": "Tercatat di lebih dari satu baris pembukuan: "
                                               + "; ".join(u["asal"] for u in us)
                                               + ". Bank dan rekeningnya tidak diisi; lengkapi lewat form."})
            continue
        u = us[0]
        # Rekening kosong di buku (41 anggota BRI) tidak menghapus rekening yang sudah ada.
        baru = {"bank": u["bank"], "rekening": u["rekening"] or p["rekening"],
                "nama_rekening": u["nama_rekening"] or p["nama_rekening"]}
        beda = {k: [p[k], baru[k]] for k in DARI_BUKU if p[k] != baru[k]}
        if beda:
            hasil["ubah"].append({"id": pid, "nama": p["nama"], "nrp": p["nrp_asli"], "asal": u["asal"],
                                  "cara": u["cara"], "perubahan": beda})
        else:
            hasil["tetap"] += 1
    terisi = {pid for pid, us in usul.items() if len(us) == 1}
    hasil["tetap_tanpa_bank"] = [{"id": p["id"], "nama": p["nama"], "nrp": p["nrp_asli"], "jenis": p["jenis"]}
                                 for p in semua if p["status"] == "aktif" and not p["bank"] and p["id"] not in terisi]
    return hasil


def terapkan(rencana: dict, operator: str) -> dict:
    """Terapkan rencana impor dalam satu transaksi. Ditolak bila database berubah sejak pratinjau."""
    with _tulis, _sambung() as con:
        if _versi(con) != rencana["versi"]:
            return {"ok": False, "galat": {"_": "Data personel berubah sejak pratinjau dibuat (mungkin di tab "
                                                "lain). Buat pratinjau ulang, lalu terapkan lagi."}}
        now, n = _sekarang(), {"tambah": 0, "ubah": 0, "tandai": 0}
        if rencana["jenis_impor"] == "nominatif":
            for b in rencana["tambah"]:
                v = {k: "" for k in KOLOM}
                v.update(b, status="aktif")
                cur = con.execute(f"INSERT INTO personel ({', '.join(KOLOM)}, dibuat, diubah) VALUES "
                                  f"({', '.join('?' * len(KOLOM))}, ?, ?)", [v[k] for k in KOLOM] + [now, now])
                _catat(con, "impor: tambah", operator, cur.lastrowid, v["nrp"], v["nama"],
                       {k: ["", v[k]] for k in DARI_NOMINATIF if v[k]})
                n["tambah"] += 1
            for x in rencana["hilang"]:
                lama = con.execute("SELECT tandai FROM personel WHERE id = ?", (x["id"],)).fetchone()
                if lama is None or lama["tandai"] == x["tandai"]:
                    continue
                con.execute("UPDATE personel SET tandai = ?, diubah = ? WHERE id = ?", (x["tandai"], now, x["id"]))
                _catat(con, "impor: ditandai", operator, x["id"], digit(x["nrp"]), x["nama"],
                       {"tandai": [lama["tandai"], x["tandai"]]})
                n["tandai"] += 1
        for x in rencana["ubah"]:
            sets = {k: b for k, (_, b) in x["perubahan"].items()}
            if "nrp_asli" in sets:
                sets["nrp"] = digit(sets["nrp_asli"])
            con.execute(f"UPDATE personel SET {', '.join(k + ' = ?' for k in sets)}, diubah = ? WHERE id = ?",
                        list(sets.values()) + [now, x["id"]])
            _catat(con, f"impor: {'bank' if rencana['jenis_impor'] == 'bank' else 'ubah'}", operator, x["id"],
                   digit(x["nrp"]), x["nama"], x["perubahan"])
            n["ubah"] += 1
        _catat(con, "impor", operator, None, "", "",
               {"berkas": [b["file"] for b in rencana["berkas"]], "jenis": rencana["jenis_impor"], **n})
        _naik(con)
    return {"ok": True, **n}
